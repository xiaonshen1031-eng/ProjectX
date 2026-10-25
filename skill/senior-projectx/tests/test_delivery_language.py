"""Current-file language enforcement inside the real Project delivery audit."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import audit_project_delivery as delivery
import check_writing_quality as writing


class ProjectDeliveryLanguageTests(unittest.TestCase):
    def setUp(self):
        # This suite isolates prose enforcement; real figure-policy integration
        # and data binding are exercised by test_rich_project_delivery.
        figure_policy = patch.object(delivery, 'audit_plot_policy', return_value={'status': 'pass', 'errors': []})
        figure_policy.start(); self.addCleanup(figure_policy.stop)
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.docx = self.root / "report.docx"
        self.write_docx("容量约束限定各时段的供能上限，能量平衡将分配结果与总充电需求关联。")
        self.pdf = self.root / "report.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        with self.pdf.open("wb") as stream:
            writer.write(stream)
        screenshot = self.root / "page.png"
        Image.new("RGB", (10, 10), "white").save(screenshot)
        self.args = argparse.Namespace(
            manifest=self.record("manifest.json", {"document_type": "project", "manuscript_mode": "project_word",
                "language": "chinese", "equation_backend": "office_math"}),
            data_audit=[self.record("data.json", {"status": "pass"})],
            result=self.record("result.json", {"status": "optimal"}),
            validation=self.record("validation.json", {"status": "pass"}),
            figures=self.record("figures.json", {"artifacts": []}),
            docx=self.docx, pdf=self.pdf, mathtype_evidence=None,
            visual_review=self.record("visual.json", {"pdf_sha256": hashlib.sha256(self.pdf.read_bytes()).hexdigest(),
                "reviewer": "unit-test-fixture", "reviewed_pages": [1], "page_images": [str(screenshot)], "status": "pass"}),
        )

    def tearDown(self):
        self.temp.cleanup()

    def record(self, name, value):
        path = self.root / name
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return path

    def write_docx(self, text):
        document = Document()
        document.add_paragraph(text)
        math = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        content = OxmlElement("m:t")
        content.text = "x=1"
        run.append(content)
        math.append(run)
        document.add_paragraph()._p.append(math)
        document.save(self.docx)

    def test_clean_current_document_passes_without_a_language_flag(self):
        result = delivery.audit(self.args)
        self.assertEqual(result["status"], "pass", result["errors"])
        self.assertEqual(result["writing_audit"]["manuscript_editorial_status"], "clear")

    def test_default_delivery_rejects_confirmed_report_language(self):
        self.write_docx("作者确认校验通过，文件尚未回传。")
        before = self.docx.read_bytes()
        result = delivery.audit(self.args)
        self.assertEqual(result["status"], "fail")
        self.assertGreater(result["writing_audit"]["unresolved_editorial_count"], 0)
        self.assertEqual(self.docx.read_bytes(), before)

    def test_old_clean_result_cannot_approve_changed_current_prose(self):
        self.assertEqual(delivery.audit(self.args)["status"], "pass")
        self.record("old_writing.json", {"manuscript_editorial_status": "clear", "unresolved_editorial_count": 0})
        self.write_docx("首先构建证据链，全面提升模型鲁棒性。")
        self.assertEqual(delivery.audit(self.args)["status"], "fail")

    def test_concrete_scientific_information_constraint_can_be_reviewed(self):
        self.write_docx("The model cannot update before a forecast error matures.")
        result = writing.inspect_writing(self.docx, manuscript_mode="project", language="chinese")
        self.args.review_decisions = self.record("review.json", {"reviews": [
            {"candidate_id": c["candidate_id"], "decision": "scientific_content",
             "reason": "The sentence defines the information constraint: forecast error maturity determines the model update time."}
            for c in result["editorial_candidates"]]})
        self.assertEqual(delivery.audit(self.args)["status"], "pass")

    def test_generic_approval_and_invalid_review_file_do_not_pass(self):
        self.write_docx("The framework cannot support operational decisions.")
        result = writing.inspect_writing(self.docx, manuscript_mode="project", language="chinese")
        self.args.review_decisions = self.record("review.json", {"reviews": [
            {"candidate_id": c["candidate_id"], "decision": "scientific_content", "reason": "Reviewed and approved."}
            for c in result["editorial_candidates"]]})
        self.assertEqual(delivery.audit(self.args)["status"], "fail")
        self.args.review_decisions.write_text("{", encoding="utf-8")
        result = delivery.audit(self.args)
        self.assertEqual(result["status"], "fail")
        self.assertTrue(any("language review" in error for error in result["errors"]))


if __name__ == "__main__":
    unittest.main()
