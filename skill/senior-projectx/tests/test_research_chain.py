from __future__ import annotations
import hashlib, json, sys, tempfile, unittest
from pathlib import Path
import pandas as pd
from docx import Document
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import audit_docx_submission, audit_project_delivery, check_writing_quality as writing
import project_report, project_research as project


class ProjectDocumentBackendTests(unittest.TestCase):
    def test_office_math_formula_is_editable_omml(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "office_math.docx"
            document = Document(); formulas = []
            project_report._add_formula(
                document, r"\sum_{t\in\mathcal{T}}x_t=E", "energy", formulas,
                equation_backend="office_math",
            )
            document.save(path)
            result = audit_docx_submission.audit(
                path, False, [], equation_backend="office_math", equations_required=True
            )
            self.assertEqual(result["equation_backend_detected"], "office_math")
            self.assertEqual(result["office_math_object_count"], 1)
            self.assertEqual(result["toggle_tex_formula_count"], 0)
            self.assertEqual(formulas[0]["formula_id"], "energy")

    def test_visual_review_is_bound_to_current_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "report.pdf"; pdf.write_bytes(b"current-pdf")
            digest = hashlib.sha256(pdf.read_bytes()).hexdigest()
            self.assertTrue(audit_project_delivery.visual_review_matches_pdf({"pdf_sha256": digest}, pdf))
            pdf.write_bytes(b"changed-pdf")
            self.assertFalse(audit_project_delivery.visual_review_matches_pdf({"pdf_sha256": digest}, pdf))

class DataSolveAndLineageTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.periods = self.root / "periods.csv"; self.vehicles = self.root / "vehicles.csv"; self.config = self.root / "config.json"
        pd.DataFrame({"period": [0, 1, 2], "base_load_kw": [0.0, 3.0, 2.0], "base_load_unit": ["kW"]*3,
                      "capacity_kw": [8.0, 8.0, 8.0], "duration_h": [1.0]*3}).to_csv(self.periods, index=False)
        pd.DataFrame({"vehicle_id": ["A", "B"], "arrival_period": [0, 1], "departure_period": [2, 3],
                      "energy_kwh": [4.0, 4.0], "max_charge_kw": [4.0, 4.0]}).to_csv(self.vehicles, index=False)
        self.config.write_text(json.dumps({"period_cost": [1, 2, 3], "validation_tolerance": 1e-7}), encoding="utf-8")
    def tearDown(self): self.tmp.cleanup()

    def test_data_audit_preserves_legal_zero_and_catches_unit_conflict(self):
        schema = self.root / "schema.json"
        schema.write_text(json.dumps({"fields": [
            {"name": "period", "role": "record_id", "unit": "index"},
            {"name": "base_load_kw", "unit": "kW", "unit_field": "base_load_unit", "allow_zero": True, "min": 0},
            {"name": "capacity_kw", "unit": "kW", "min": 0},
        ]}), encoding="utf-8")
        passed = project.audit_data(self.periods, schema)
        self.assertEqual(passed["status"], "pass")
        self.assertEqual(passed["profiles"]["base_load_kw"]["zero_count"], 1)
        frame = pd.read_csv(self.periods); frame.loc[2, "base_load_unit"] = "MW"; frame.to_csv(self.periods, index=False)
        failed = project.audit_data(self.periods, schema)
        self.assertIn("unit_conflict", {e["type"] for e in failed["errors"]})

    def test_continuous_lp_and_independent_validation(self):
        allocation = self.root / "allocation.csv"
        result = project.solve_allocation(self.periods, self.vehicles, self.config, allocation)
        self.assertEqual(result["status"], "optimal")
        result_path = self.root / "result.json"; project.write_json(result_path, result)
        validation = project.validate_allocation(self.periods, self.vehicles, self.config, allocation, result_path)
        self.assertEqual(validation["status"], "pass")
        self.assertLess(validation["maximum_energy_residual_kwh"], 1e-7)
        bad = json.loads(self.config.read_text()); bad["integer_required"] = True; self.config.write_text(json.dumps(bad), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "continuous LPs only"):
            project.solve_allocation(self.periods, self.vehicles, self.config, self.root / "bad.csv")

    def test_structured_solver_boundaries(self):
        solved = project.bounded_sum_linear([2, 1], [0, 0], [2, 3], 4)
        self.assertEqual(solved["status"], "optimal"); self.assertAlmostEqual(sum(solved["x"]), 4)
        replay = project.dag_timing_replay([
            {"id": "a", "duration": 2, "resource": "r"},
            {"id": "b", "duration": 1, "resource": "r", "depends_on": ["a"]}], ["a", "b"])
        self.assertEqual(replay["scope"], "given order only")
        with self.assertRaises(ValueError):
            project.dag_timing_replay([{"id": "a", "duration": 2}, {"id": "b", "duration": 1, "depends_on": ["a"]}], ["b", "a"])

    def test_lineage_invalidates_only_declared_consumers(self):
        manifest = self.root / "project_manifest.json"; manifest.write_text(json.dumps({"project": {"lineage": []}}), encoding="utf-8")
        artifact = self.root / "figure.png"; artifact.write_bytes(b"figure")
        project.lineage_record(manifest, artifact, [self.periods], role="figure")
        self.assertEqual(project.lineage_check(manifest)["status"], "pass")
        self.periods.write_text(self.periods.read_text() + "\n", encoding="utf-8")
        checked = project.lineage_check(manifest)
        self.assertEqual(checked["artifacts"][0]["status"], "stale")

class BilingualWritingTests(unittest.TestCase):
    def run_doc(self, paragraphs, language="auto", mode="project"):
        path = self.root / "draft.docx"; doc = Document()
        for text in paragraphs: doc.add_paragraph(text)
        doc.save(path)
        return writing.inspect_writing(path, manuscript_mode=mode, language=language)
    def setUp(self): self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
    def tearDown(self): self.tmp.cleanup()

    def test_chinese_risks_have_contextual_diagnostics(self):
        result = self.run_doc(["1 研究结果", "根据提示，本次生成全面系统地分析了结果。",
                               "本研究提出“全域协同增益机制”，强调稳定性、可靠性、鲁棒性和泛化性——同时提升——全面增强。",
                               "[待补]"])
        kinds = {w["type"] for w in result["warnings"]}
        self.assertTrue({"process_language_leakage", "invented_quote_concept", "abstract_noun_stack", "em_dash_chain", "definite_placeholder"} <= kinds)
        self.assertEqual(result["language"], "chinese")
        self.assertGreater(result["blocking_warning_count"], 0)
        for warning in result["warnings"]:
            self.assertIn("reason", warning); self.assertIn("suggestion", warning); self.assertIn("original_text", warning)

    def test_legitimate_terms_minus_ranges_and_real_risk_survive(self):
        result = self.run_doc(["1 结果", "字段“负荷率”沿用任务书定义。温度范围为−5–10 ℃。容量不足会导致方案不可行，因此保留该工程风险。"], "chinese")
        kinds = {w["type"] for w in result["warnings"]}
        self.assertFalse(kinds & {"invented_quote_concept", "em_dash_chain", "hollow_report_prose", "definite_placeholder"})

    def test_english_branch_has_no_chinese_length_false_positive(self):
        result = self.run_doc(["3 Results", "The allocation satisfies every energy balance and capacity constraint under the stated operating case."], "english", "journal")
        self.assertEqual(result["language"], "english")
        self.assertNotIn("chinese_repetitive_sentence_skeleton", {w["type"] for w in result["warnings"]})
