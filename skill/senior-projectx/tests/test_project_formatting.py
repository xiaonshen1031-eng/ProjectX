"""Project-only regressions for connected prose, typography and equation sources."""
from __future__ import annotations

import base64
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import project_report as report


class ProjectFormattingTests(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        report._configure_styles(self.doc)
        report._configure_page(self.doc)

    def test_default_fonts_are_black_regular_and_theme_free(self):
        for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
            style = self.doc.styles[name]
            self.assertFalse(style.font.bold, name)
            self.assertEqual(str(style.font.color.rgb), "000000", name)
            fonts = style._element.rPr.rFonts
            self.assertEqual(fonts.get(qn("w:eastAsia")), "SimSun")
            self.assertEqual(fonts.get(qn("w:ascii")), "Times New Roman")
            self.assertFalse(any("Theme" in key for key in fonts.attrib))
        self.assertEqual(self.doc.styles["Normal"].font.size.pt, 12)

    def test_body_cross_reference_is_not_a_caption(self):
        body = report._add_body(self.doc, "图1与表3共同给出容量余量及其变化。")
        caption = report._add_caption(self.doc, "图1 容量余量")
        self.assertEqual(body.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertEqual(body._p.pPr.ind.get(qn("w:firstLineChars")), "200")
        self.assertEqual(body.paragraph_format.first_line_indent.pt, 24)
        spacing = body._p.pPr.spacing
        self.assertEqual((spacing.get(qn("w:line")), spacing.get(qn("w:lineRule"))), ("360", "auto"))
        self.assertEqual(caption.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(caption._p.pPr.ind.get(qn("w:firstLineChars")), "0")

    def test_three_line_table_has_no_grid_or_shading_and_no_spacer(self):
        before = len(self.doc.paragraphs)
        table = report._add_table(self.doc, ["对象", "数值"], [["A", "0.00"], ["B", "1.00"]], [8, 8])
        self.assertEqual(len(self.doc.paragraphs), before)
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        for edge in ["top", "bottom"]:
            node = borders.find(qn(f"w:{edge}"))
            self.assertEqual((node.get(qn("w:val")), node.get(qn("w:color"))), ("single", "000000"))
        for edge in ["left", "right", "insideH", "insideV"]:
            self.assertEqual(borders.find(qn(f"w:{edge}")).get(qn("w:val")), "nil")
        for index, row in enumerate(table.rows):
            for cell in row.cells:
                tcpr = cell._tc.tcPr
                shading = tcpr.find(qn("w:shd"))
                self.assertEqual((shading.get(qn("w:val")), shading.get(qn("w:fill"))), ("clear", "auto"))
                if index == 0:
                    self.assertEqual(tcpr.find(qn("w:tcBorders")).find(qn("w:bottom")).get(qn("w:color")), "000000")
                p = cell.paragraphs[0]
                self.assertEqual(p._p.pPr.ind.get(qn("w:firstLineChars")), "0")
                self.assertEqual(p._p.pPr.spacing.get(qn("w:line")), "240")
                self.assertEqual(p._p.pPr.spacing.get(qn("w:lineRule")), "auto")
                for run in p.runs:
                    self.assertEqual(run.font.size.pt, 10.5)
                    self.assertFalse(run.font.bold)
                    self.assertEqual(str(run.font.color.rgb), "000000")
        self.assertLessEqual(sum(c.width for c in table.columns), self.doc.sections[0].page_width -
                             self.doc.sections[0].left_margin - self.doc.sections[0].right_margin + Pt(1))

    def test_display_and_reference_roles_reset_inherited_layout(self):
        formulas = []
        p = report._add_formula(self.doc, "H=C-L", "headroom", formulas)
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(p._p.pPr.ind.get(qn("w:firstLineChars")), "0")
        self.assertEqual(p.paragraph_format.first_line_indent.pt, 0)
        ref = self.doc.add_paragraph("Author A. A long English title.")
        report._paragraph_role(ref, "reference")
        self.assertEqual(ref.alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_inline_inventory_keeps_case_repetition_order_and_table_size(self):
        formulas = []
        report._add_body(self.doc, "容量 $C$ 与成本 $c$ 不同，容量仍记为 $C$。", formulas=formulas)
        report._add_formula(self.doc, "H=C-L", "headroom", formulas)
        report._add_table(self.doc, ["容量 $C$", "功率"], [["0", "$x_t$"]], formulas=formulas)
        self.assertEqual([f["tex"] for f in formulas], ["C", "c", "C", "H=C-L", "C", "x_t"])
        self.assertEqual([f["placement"] for f in formulas], ["inline"] * 3 + ["display", "inline", "inline"])
        self.assertEqual([f["context"] for f in formulas][-2:], ["table", "table"])
        self.assertEqual([f["font_size_pt"] for f in formulas][-2:], [10.5, 10.5])
        self.assertEqual(len({f["formula_id"] for f in formulas}), len(formulas))
        with self.assertRaisesRegex(ValueError, "manifest"):
            report._add_body(self.doc, "功率 $x_t$。")

    def test_inline_office_math_uses_real_objects_at_context_size(self):
        try:
            report._office_math_xslt()
        except RuntimeError as exc:
            self.skipTest(str(exc))
        formulas = []
        p = report._add_body(self.doc, "容量 $C_t$。", formulas=formulas, equation_backend="office_math")
        table = report._add_table(self.doc, ["功率 $x_t$"], [["1.0"]], formulas=formulas, equation_backend="office_math")
        self.assertEqual(len(p._p.xpath(".//m:oMath")), 1)
        math = table.cell(0, 0)._tc.xpath(".//m:oMath")
        self.assertEqual(len(math), 1)
        # lxml's XSLT result keeps MathML namespaces; register Word explicitly.
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        self.assertTrue(all(n.get(qn("w:val")) == "21" for n in math[0].xpath(".//w:sz", namespaces=ns)))

    def test_build_example_has_no_internal_paths_or_verification_narration(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            def save(name, payload):
                path = root / name
                path.write_text(json.dumps(payload), encoding="utf-8")
                return path
            template = root / "template.docx"; report.create_default_template(template)
            manifest = save("manifest.json", {"paper_title": "充电计划分析", "equation_backend": "mathtype_native"})
            result = save("result.json", {"status": "optimal", "objective": 2.5, "capacity_margin_kw": [1, 2]})
            validation = save("validation.json", {"status": "pass", "minimum_capacity_margin_kw": 1,
                              "maximum_energy_residual_kwh": 0, "maximum_capacity_violation_kw": 0})
            brief = save("brief.json", {"task_id": "T1", "private_path": "PRIVATE_INTERNAL_PATH"})
            audit = save("audit.json", {"status": "pass", "rows": 2, "data_path": "PRIVATE_INTERNAL_PATH",
                                       "profiles": {"base_load_kw": {"unit": "kW"}}})
            pic = root / "pixel.png"
            pic.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j8ioAAAAASUVORK5CYII="))
            figures = save("figures.json", {"artifacts": [{"id": ident, "path": str(pic), "width_cm": 2,
                            "caption": f"图{i} 调度结果"} for i, ident in enumerate(
                            ["figure_process", "figure_balance", "figure_sensitivity"], 1)]})
            sensitivity = root / "sensitivity.csv"
            sensitivity.write_text("demand_scale,status,objective,mean_margin_kw\n1.0,optimal,2.5,1.5\n", encoding="utf-8")
            output = root / "source.docx"; inventory = root / "source.json"
            built = report.build_report(template, manifest, [audit], result, validation, brief,
                                        figures, sensitivity, output, inventory)
            doc = Document(output)
            text = "\n".join(p.text for p in doc.paragraphs)
            for token in ["PRIVATE_INTERNAL_PATH", "JSON 指针", "章节 brief", "证据链", "success 布尔", "复核职责"]:
                self.assertNotIn(token, text)
            self.assertIn("合成", text)  # Do not erase the factual source classification.
            self.assertEqual(built["status"], "conversion_pending")
            source = json.loads(inventory.read_text(encoding="utf-8"))
            self.assertTrue(any(f["placement"] == "inline" for f in source["formulas"]))
            self.assertEqual(source["formula_count"], len(source["formulas"]))
            self.assertFalse(source["publication_eligible"])


class PowerShellTokenTests(unittest.TestCase):
    def test_actual_converter_inventory_is_case_sensitive(self):
        shell = shutil.which("pwsh") or shutil.which("powershell")
        if not shell:
            self.skipTest("PowerShell unavailable")
        # Parse the real converter and execute only its pure inventory function: no Word launch.
        script = r'''
$errors = $null; $tokens = $null
$ast = [System.Management.Automation.Language.Parser]::ParseFile($env:PROJECTX_TOKEN_TEST_SCRIPT, [ref]$tokens, [ref]$errors)
if ($errors.Count) { throw ($errors | Out-String) }
foreach ($name in @('Get-ExpectedFormulaTokens', 'Assert-SourceFormulaTokens')) {
    $fn = $ast.Find({ param($n) $n -is [System.Management.Automation.Language.FunctionDefinitionAst] -and $n.Name -eq $name }, $true)
    if ($null -eq $fn) { throw 'Missing inventory function' }
    Invoke-Expression $fn.Extent.Text
}
$formulas = @(
    [pscustomobject]@{placement='inline';tex='C'},
    [pscustomobject]@{placement='inline';tex='c'},
    [pscustomobject]@{placement='inline';tex='C'},
    [pscustomobject]@{placement='display';tex='C=c'}
)
$counts = Get-ExpectedFormulaTokens $formulas
if ($counts.Count -ne 3 -or $counts['$C$'] -ne 2 -or $counts['$c$'] -ne 1 -or $counts['\[C=c\]'] -ne 1) { throw 'Case/repetition inventory mismatch' }
Assert-SourceFormulaTokens 'Capacity $C$, cost $c$, repeated $C$. \[C=c\]' $formulas
foreach ($bad in @('$c$ $C$ $C$ \[C=c\]', '$C$ $C$ $C$ \[C=c\]', '$C$ $c$ $C$ \[C=c\] $z$')) {
    $rejected = $false
    try { Assert-SourceFormulaTokens $bad $formulas } catch { $rejected = $true }
    if (-not $rejected) { throw 'Invalid source inventory was accepted' }
}
'case-sensitive inventory passed'
'''
        env = dict(os.environ, PROJECTX_TOKEN_TEST_SCRIPT=str(ROOT / "scripts/project_mathtype_convert.ps1"))
        run = subprocess.run([shell, "-NoProfile", "-NonInteractive", "-Command", script], env=env,
                             capture_output=True, text=True, timeout=30)
        self.assertEqual(run.returncode, 0, run.stdout + run.stderr)
        self.assertIn("passed", run.stdout)


if __name__ == "__main__":
    unittest.main()
