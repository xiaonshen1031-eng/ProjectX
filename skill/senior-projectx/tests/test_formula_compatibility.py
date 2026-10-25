from __future__ import annotations
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'scripts'))
from build_formula_manifest import build_manifest


class FormulaCompatibilityTests(unittest.TestCase):
    def test_inventory_preserves_case_occurrences_and_table_context(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            doc = Document()
            doc.add_paragraph(r'容量 $C$ 与成本 $c$ 分别定义。')
            doc.add_paragraph(r'\[C=c+1\]')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = r'$C$'
            path = root / 'source.docx'
            doc.save(path)
            result = build_manifest(path, root / 'source.json')
            self.assertEqual([f['tex'] for f in result['formulas']], ['C', 'c', 'C=c+1', 'C'])
            self.assertEqual(result['formulas'][-1]['context'], 'table')
            self.assertEqual(result['formulas'][2]['placement'], 'display')
            self.assertEqual(result['schema'], 'senior_projectx.project.equation_source.v1')

    def test_native_source_is_not_toggled_again(self):
        from docx.oxml import OxmlElement
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            doc = Document()
            doc.add_paragraph()._p.append(OxmlElement('m:oMath'))
            doc.save(root / 'native.docx')
            with self.assertRaisesRegex(ValueError, 'Existing OLE/OMML'):
                build_manifest(root / 'native.docx', root / 'source.json')

    def test_actual_schema_guard_accepts_both_project_identities_and_rejects_unknown(self):
        shell = shutil.which('powershell') or shutil.which('pwsh')
        if not shell:
            self.skipTest('PowerShell is required for the conversion guard.')
        script = r'''
$errors = $null; $tokens = $null
$ast = [System.Management.Automation.Language.Parser]::ParseFile($env:PROJECTX_CONVERTER, [ref]$tokens, [ref]$errors)
if ($errors.Count) { throw ($errors | Out-String) }
foreach ($name in @('Assert-FormulaSourceSchema', 'Get-RelativeFilePath')) {
    $fn = $ast.Find({ param($n) $n -is [System.Management.Automation.Language.FunctionDefinitionAst] -and $n.Name -eq $name }, $true)
    if ($null -eq $fn) { throw 'Required function was not found.' }
    Invoke-Expression $fn.Extent.Text
}
Assert-FormulaSourceSchema 'senior_projectx.project.equation_source.v1'
Assert-FormulaSourceSchema 'nps.project.equation_source.v1'
foreach ($schema in @('unknown.v1', 'nps.project.equation_source.v2', 'NPS.PROJECT.EQUATION_SOURCE.V1')) {
    $rejected = $false
    try { Assert-FormulaSourceSchema $schema } catch { $rejected = $true }
    if (-not $rejected) { throw 'An unrelated formula schema was accepted.' }
}
$relative = Get-RelativeFilePath 'C:\task\review' 'C:\task\source.docx'
if ($relative -ne '..\source.docx') { throw ('Unexpected relative path: ' + $relative) }
'compatible guards passed'
'''
        ran = subprocess.run([shell, '-NoProfile', '-NonInteractive', '-Command', script], env={**os.environ, 'PROJECTX_CONVERTER': str(ROOT / 'scripts/project_mathtype_convert.ps1')}, capture_output=True, text=True, timeout=30)
        self.assertEqual(ran.returncode, 0, ran.stdout + ran.stderr)
        self.assertIn('compatible guards passed', ran.stdout)
