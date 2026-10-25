"""Behavioral regressions for Project-only numbering and plotting typography.

OLE fixtures deliberately test preservation/layout, not native authenticity.
"""
from __future__ import annotations

import importlib.util
from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import project_equation_layout as layout
import project_report as report


class EquationLayoutTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def fixture(self, *, text="", columns=1, placement="body", extra=None):
        doc = Document()
        display = doc.add_paragraph()
        display.style = doc.styles["Normal"]
        props = display._p.get_or_add_pPr()
        props.append(layout.element("w:tabs"))
        props.find(layout.tag("w:tabs")).append(layout.element("w:tab", val="center", pos=4100))
        style_tabs = layout.element("w:tabs")
        style_tabs.append(layout.element("w:tab", val="right", pos=9800))
        doc.styles["Normal"]._element.get_or_add_pPr().append(style_tabs)
        props.append(layout.element("w:ind", left=200, firstLine=480, firstLineChars=200))
        start = layout.element("w:bookmarkStart", id=23, name="formula_anchor")
        end = layout.element("w:bookmarkEnd", id=23)
        display._p.append(start)
        formula_run = display.add_run()._r
        obj = layout.element("w:object")
        ole = layout.element("o:OLEObject")
        ole.set("ProgID", "Equation.DSMT4")
        ole.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "rIdEquationTest")
        obj.append(ole); formula_run.append(obj)
        display._p.append(end)
        if text:
            display.add_run(text)
        if extra is not None:
            display._p.append(extra)
        if placement == "table":
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0)._tc.append(display._p)
        inline = doc.add_paragraph("设备容量为")
        inline_run = inline.add_run()._r
        inline_obj = layout.element("w:object")
        inline_ole = layout.element("o:OLEObject")
        inline_ole.set("ProgID", "Equation.DSMT4")
        inline_obj.append(inline_ole); inline_run.append(inline_obj)
        inline.add_run("。")
        section = doc.sections[0]._sectPr
        page = section.find("w:pgSz", layout.NS); page.set(layout.tag("w:w"), "12240")
        margins = section.find("w:pgMar", layout.NS)
        for key, value in [("left", 1200), ("right", 1440), ("gutter", 200)]:
            margins.set(layout.tag("w:" + key), str(value))
        section.find("w:cols", layout.NS).set(layout.tag("w:num"), str(columns))
        source = self.root / "source.docx"
        doc.save(source)
        # Synthetic payloads are NOT represented as valid Equation Native streams.
        with ZipFile(source, "a", compression=ZIP_DEFLATED) as archive:
            archive.writestr("word/embeddings/layout_fixture.bin", b"synthetic-preservation-fixture\x00\x01")
            archive.writestr("word/media/layout_preview.emf", b"synthetic-preview-fixture")
        manifest = self.root / "inventory.json"
        manifest.write_text(json.dumps({"formula_count": 2, "formulas": [
            {"formula_id": "display_A", "placement": "display"},
            {"formula_id": "inline_A", "placement": "inline"}]}), encoding="utf-8")
        return source, manifest

    def parse(self, path):
        with ZipFile(path) as archive:
            return etree.fromstring(archive.read("word/document.xml"))

    def test_center_and_right_tabs_follow_current_width_and_number_is_a_field(self):
        source, manifest = self.fixture()
        output = self.root / "numbered.docx"
        result = layout.number_docx(source, manifest, output)
        p = self.parse(output).find("w:body/w:p", layout.NS)
        tabs = p.findall("w:pPr/w:tabs/w:tab", layout.NS)
        self.assertEqual([(t.get(layout.tag("w:val")), t.get(layout.tag("w:pos"))) for t in tabs],
                         [("clear", "4100"), ("clear", "9800"), ("center", "4700"), ("right", "9400")])
        self.assertEqual(p.find("w:pPr/w:jc", layout.NS).get(layout.tag("w:val")), "left")
        self.assertTrue(all(v == "0" for v in p.find("w:pPr/w:ind", layout.NS).attrib.values()))
        field = p.find("w:fldSimple", layout.NS)
        self.assertIn("SEQ Equation", field.get(layout.tag("w:instr")))
        self.assertEqual(field.find("w:r/w:t", layout.NS).text, "1")
        self.assertEqual(field.find("w:r/w:rPr/w:sz", layout.NS).get(layout.tag("w:val")), "24")
        self.assertEqual(field.find("w:r/w:rPr/w:i", layout.NS).get(layout.tag("w:val")), "0")
        self.assertEqual(result["display_equations_numbered"], 1)
        self.assertEqual(result["inline_equations_unchanged"], 1)
        self.assertEqual(result["status"], "layout_applied_review_pending")

    def test_native_runs_previews_all_other_parts_and_inline_paragraph_are_unchanged(self):
        source, manifest = self.fixture()
        output = self.root / "numbered.docx"
        before = self.parse(source)
        layout.number_docx(source, manifest, output)
        after = self.parse(output)
        native = lambda root: [etree.tostring(n) for n in root.xpath(".//w:r[w:object]", namespaces=layout.NS)]
        self.assertEqual(native(before), native(after))
        self.assertEqual(etree.tostring(before.find("w:body/w:p[2]", layout.NS)),
                         etree.tostring(after.find("w:body/w:p[2]", layout.NS)))
        for kind in ["w:bookmarkStart", "w:bookmarkEnd"]:
            self.assertEqual([etree.tostring(n) for n in before.findall(".//" + kind, layout.NS)],
                             [etree.tostring(n) for n in after.findall(".//" + kind, layout.NS)])
        with ZipFile(source) as original, ZipFile(output) as final:
            self.assertEqual(original.namelist(), final.namelist())
            for name in original.namelist():
                if name != "word/document.xml":
                    self.assertEqual(original.read(name), final.read(name), name)

    def test_repeat_layout_is_idempotent_and_never_overwrites_source(self):
        source, manifest = self.fixture()
        first, second = self.root / "first.docx", self.root / "second.docx"
        original = source.read_bytes()
        layout.number_docx(source, manifest, first)
        layout.number_docx(first, manifest, second)
        self.assertEqual(etree.tostring(self.parse(first)), etree.tostring(self.parse(second)))
        self.assertEqual(source.read_bytes(), original)
        with self.assertRaises(FileExistsError):
            layout.number_docx(source, manifest, source)
        with self.assertRaises(FileExistsError):
            layout.number_docx(source, manifest, first)

    def test_ambiguous_prose_table_and_multicolumn_fail_before_creating_output(self):
        for kwargs, message in [({"text": "这里是正文，不应删除"}, "Text around"),
                                ({"columns": 2}, "Multi-column"), ({"placement": "table"}, "ordinary body")]:
            source, manifest = self.fixture(**kwargs)
            output = self.root / "must_not_exist.docx"
            with self.assertRaisesRegex(ValueError, message):
                layout.number_docx(source, manifest, output)
            self.assertFalse(output.exists())

    def test_unrelated_fields_and_footnotes_are_not_deleted(self):
        footnote = layout.element("w:r"); footnote.append(layout.element("w:footnoteReference", id=1))
        ref = layout.element("w:fldSimple", instr=" REF old_target ")
        for extra in [ref, footnote]:
            source, manifest = self.fixture(extra=extra)
            output = self.root / "must_not_exist.docx"
            with self.assertRaises(ValueError):
                layout.number_docx(source, manifest, output)
            self.assertFalse(output.exists())

    def test_section_lookup_uses_following_break_not_last_section(self):
        root = etree.fromstring(f'<w:body xmlns:w="{layout.NS["w"]}"><w:p/><w:p><w:pPr><w:sectPr>'
            '<w:pgSz w:w="10000"/><w:pgMar w:left="1000" w:right="1000"/>'
            '</w:sectPr></w:pPr></w:p><w:sectPr><w:pgSz w:w="14000"/>'
            '<w:pgMar w:left="1000" w:right="1000"/></w:sectPr></w:body>')
        self.assertEqual(layout.section_width(root[0], root), 8000)
        self.assertEqual(layout.section_width(root[0], root, gutter_at_top=True), 8000)

    def test_inline_before_display_does_not_consume_a_number(self):
        source, manifest = self.fixture()
        root = self.parse(source)
        body = root.find("w:body", layout.NS)
        inline = body[1]; body.remove(inline); body.insert(0, inline)
        changed = self.root / "inline_first.docx"
        with ZipFile(source) as archive, ZipFile(changed, "w") as final:
            for info in archive.infolist():
                final.writestr(info, etree.tostring(root) if info.filename == "word/document.xml" else archive.read(info.filename))
        record = json.loads(manifest.read_text())
        record["formulas"].reverse(); manifest.write_text(json.dumps(record), encoding="utf-8")
        output = self.root / "numbered.docx"
        layout.number_docx(changed, manifest, output)
        fields = self.parse(output).xpath(".//w:fldSimple/w:r/w:t/text()", namespaces=layout.NS)
        self.assertEqual(fields, ["1"])

    def test_multiple_displays_have_continuous_cached_fields_on_repeat(self):
        source, manifest = self.fixture()
        root = self.parse(source)
        body = root.find("w:body", layout.NS)
        second_display = deepcopy(body[0])
        for bookmark in second_display.findall("w:bookmarkStart", layout.NS) + second_display.findall("w:bookmarkEnd", layout.NS):
            second_display.remove(bookmark)
        body.insert(len(body) - 1, second_display)
        changed = self.root / "two_displays.docx"
        with ZipFile(source) as archive, ZipFile(changed, "w") as final:
            for info in archive.infolist():
                final.writestr(info, etree.tostring(root) if info.filename == "word/document.xml" else archive.read(info.filename))
        record = json.loads(manifest.read_text())
        record["formula_count"] = 3
        record["formulas"].append({"formula_id": "display_B", "placement": "display"})
        manifest.write_text(json.dumps(record), encoding="utf-8")
        first, second = self.root / "first.docx", self.root / "second.docx"
        layout.number_docx(changed, manifest, first)
        layout.number_docx(first, manifest, second)
        for output in [first, second]:
            self.assertEqual(self.parse(output).xpath(".//w:fldSimple/w:r/w:t/text()", namespaces=layout.NS), ["1", "2"])
        self.assertEqual(etree.tostring(self.parse(first)), etree.tostring(self.parse(second)))

    def test_existing_number_bookmark_is_not_silently_detached(self):
        source, manifest = self.fixture(text="(7)")
        root = self.parse(source)
        p = root.find("w:body/w:p", layout.NS)
        start, end = p.find("w:bookmarkStart", layout.NS), p.find("w:bookmarkEnd", layout.NS)
        p.remove(start); p.remove(end)
        label = p.findall("w:r", layout.NS)[-1]
        p.insert(p.index(label), start); p.append(end)
        changed = self.root / "number_bookmark.docx"
        with ZipFile(source) as archive, ZipFile(changed, "w") as final:
            for info in archive.infolist():
                final.writestr(info, etree.tostring(root) if info.filename == "word/document.xml" else archive.read(info.filename))
        output = self.root / "must_not_exist.docx"
        with self.assertRaisesRegex(ValueError, "cross-reference"):
            layout.number_docx(changed, manifest, output)
        self.assertFalse(output.exists())


class PlotAndContractTests(unittest.TestCase):
    def test_contract_agrees_with_effective_word_builder_defaults(self):
        contract = json.loads((ROOT / "templates/project_format_contract.json").read_text(encoding="utf-8"))
        doc = Document(); report._configure_styles(doc)
        self.assertEqual(contract["headings"]["chinese_font"], "SimSun")
        self.assertEqual(contract["body"]["size_pt"], doc.styles["Normal"].font.size.pt)
        self.assertEqual(contract["headings"]["bold"], doc.styles["Heading 1"].font.bold)
        self.assertEqual(contract["tables"]["size_pt"], 10.5)
        self.assertFalse(contract["inline_equations"]["numbered"])

    def test_builder_styles_chinese_categories_and_numeric_ticks_separately(self):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        fig, ax = plt.subplots()
        self.addCleanup(lambda: plt.close(fig))
        ax.set_xticks([0, 1], ["地区甲", "12"])
        ax.set_xlabel("时段（h）")
        fonts = {"chinese": "Microsoft YaHei", "latin": "Times New Roman", "chinese_weight": "bold"}
        report._mixed_axis(ax, fonts)
        cjk, numeric = ax.get_xticklabels()
        self.assertEqual(cjk.get_fontfamily(), ["Times New Roman", "Microsoft YaHei"])
        self.assertEqual(cjk.get_fontweight(), "bold")
        self.assertEqual(numeric.get_fontfamily(), ["Times New Roman"])
        self.assertEqual(numeric.get_fontweight(), "normal")
        self.assertEqual(ax.xaxis.label.get_fontsize(), 10)

    def test_python_starter_handles_legend_secondary_axis_and_mixed_labels(self):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        spec = importlib.util.spec_from_file_location("project_figure_starter", ROOT / "templates/python_project_figure_template.py")
        starter = importlib.util.module_from_spec(spec); spec.loader.exec_module(starter)
        with tempfile.TemporaryDirectory() as tmp:
            fig, ax = starter.new_figure()
            ax.plot([0, 1], [0, 1], label="区域 A"); ax.legend()
            ax.set_xticks([0, 1], ["地区甲", "地区乙"])
            ax2 = ax.twinx(); ax2.set_ylabel("功率（MW）")
            path = Path(tmp) / "figure.png"
            starter.finish(ax, "时段（h）", "负荷（MW）", path)
            self.assertTrue(path.is_file())
            for label in [ax.get_xticklabels()[0], ax.get_legend().get_texts()[0], ax2.yaxis.label]:
                self.assertIn(starter.CHINESE, label.get_fontfamily())
                self.assertIn(starter.LATIN, label.get_fontfamily())
            self.assertFalse(plt.fignum_exists(fig.number))


if __name__ == "__main__":
    unittest.main()
