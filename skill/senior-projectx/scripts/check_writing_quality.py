#!/usr/bin/env python3
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from difflib import SequenceMatcher
import json
import re
import sys
from pathlib import Path

from docx import Document

from manuscript_prose_rules import editorial_screen, read_reviews


PROMPT_RESIDUE = [
    "according to your request",
    "as requested",
    "after revision",
    "we revised",
    "this paper has been modified",
    "chatgpt",
    "codex",
    "according to the reviewer",
    "reviewer requested",
    "response letter",
    "changes in the manuscript",
]

OVERCLAIM_TERMS = [
    "state-of-the-art",
    "breakthrough",
    "universal",
    "guaranteed",
    "complete coverage",
    "fully solves",
    "perfect",
    "excellent performance",
    "remarkable performance",
    "robustly generalizes",
    "definitively illustrates",
    "definitively demonstrates",
    "catastrophic",
    "strictly confined",
    "robust morphological adaptability",
    "systematically corrects",
]

WEAK_PHRASES = [
    "it can be clearly seen",
    "obviously",
    "good performance",
    "greatly improves",
    "very important",
    "a lot of",
    "many aspects",
    "it is objectively noted",
    "possesses clear physical significance",
    "comprehensively summarized",
]

DEFENSIVE_PHRASES = [
    "it should be emphasized that",
    "it is worth noting that",
    "it must be noted that",
    "despite this limitation",
    "nevertheless, the proposed",
    "this does not undermine",
    "does not affect the validity",
    "remains valid despite",
    "although the result is not significant",
    "we do not claim",
    "should not be interpreted as",
    "rather than claiming",
    "does not imply", "cannot establish", "this does not mean", "despite not", "it should be noted that",
    "需要强调的是", "值得注意的是", "必须指出的是", "尽管存在上述局限",
    "这并不意味着", "这并不影响", "仍然能够证明", "不可否认",
]

CHINESE_HOLLOW_REPORT_RE = re.compile(
    r"(?:全面(?:系统)?地?(?:分析|研究|揭示)|取得(?:了)?(?:良好|显著|优异)(?:效果|成效)|"
    r"具有(?:十分|非常)?重要(?:的)?(?:意义|价值)|为.{0,24}提供(?:了)?(?:有力|坚实|重要)(?:支撑|保障)|"
    r"实现了.{0,18}(?:全面提升|显著改善))"
)
CHINESE_PROCESS_LEAKAGE_RE = re.compile(
    r"(?:根据(?:本次)?(?:提示|要求)|按照(?:用户|审稿人|上述)要求|在修改过程中|本次生成|"
    r"后续将补充|待进一步补充|由(?:ChatGPT|Codex)生成)", re.IGNORECASE
)
CHINESE_ABSTRACT_STACK_RE = re.compile(
    r"(?:[\u3400-\u9fff]{2,10}(?:性|化|能力|机制|体系|水平)[、，,]){2,}"
    r"[\u3400-\u9fff]{2,10}(?:性|化|能力|机制|体系|水平)"
)
INVENTED_QUOTE_RE = re.compile(
    r"(?:提出|构建|定义|称为|形成)(?:了)?[“\"]([^”\"]{2,18}(?:机制|范式|效应|能力|体系))[”\"]"
)
DEFINITE_PLACEHOLDER_RE = re.compile(
    r"(?:\[\s*(?:待补|待填|TODO|TBD|placeholder)\s*\]|<\s*(?:待补|TODO|TBD)\s*>|Lorem ipsum)", re.IGNORECASE
)

EXPERIMENT_GOVERNANCE_RESIDUE = [
    # Chinese project-status, provenance, and reviewer-proof residue.
    "作者确认", "作者确认的", "云端常规测试", "独立云端测试", "云端测试", "云端运行",
    "云端详细硬件清单", "尚未回传", "尚未同步", "后续归档", "逐帧原始导出",
    "只报告模型种子", "单训练种子", "单种子范围", "未估计训练不确定性",
    "不提供训练种子方差", "不能估计训练稳定性", "不能估计统计显著性",
    "本文不把", "不把该数据解释为", "不能据此声称", "尚不能作为", "不能外推",
    "尚不能外推", "这一负结果表明", "更稳妥的结论是", "现有证据仍受",
    "本文结论应限定", "当前点估计", "当前冻结仿真协议", "该图不替代",
    "而不是观察测试结果后", "未来需通过", "后续工作将优先", "尚未完成",
    "没有执行五种子", "未执行五种子", "后续归档后补充", "不进入最终质量统计",
    "不进入最终统计",
    # English equivalents.
    "author-confirmed", "confirmed by the author", "cloud test", "cloud run",
    "raw per-frame export has not been returned", "hardware inventory is incomplete",
    "single training seed", "only one training seed", "training uncertainty was not estimated",
    "cannot estimate statistical significance", "should not be interpreted as",
    "we do not interpret this as", "cannot be extrapolated", "should not be extrapolated",
    "this negative result indicates", "a more cautious conclusion is",
    "the current evidence is limited by", "the conclusion should be restricted to",
    "has not yet been synchronized",
    "has not yet been completed", "not selected after observing the test results",
    "this figure does not replace independent test evidence",
]

# Project-development lineage is useful in notes and provenance logs but is not
# a scientific explanation. Keep this list deliberately narrow so phrases such
# as "previous studies" and "prior literature" remain valid scholarly prose.
PROJECT_LINEAGE_RESIDUE = [
    "上一篇工作", "前一篇工作", "上一篇论文", "前一篇论文", "前文工作",
    "继承自上一篇", "沿用上一篇", "来自上一篇工作", "此前项目版本",
    "our previous paper", "our previous work", "the previous paper",
    "the previous work", "in our earlier paper", "in our earlier work",
    "in the prior project version", "in the preceding project version",
]

GOVERNANCE_SENSITIVE_SECTIONS = {
    "front_matter", "abstract", "introduction", "methodology", "methods", "case study",
    "experiments", "experimental settings", "results", "discussion", "limitations",
    "conclusion", "conclusions", "data and code availability",
}

CONTRAST_PHRASE_RE = re.compile(
    r"\b(?:rather than|but (?:does|do|did|is|are|was|were) not|whereas.{0,80}?remains?|"
    r"yet.{0,80}?remains?|although.{0,80}?(?:nevertheless|still)|however.{0,80}?(?:does|do|is|are) not)\b",
    re.IGNORECASE,
)
MECHANICAL_ENUMERATION_RE = re.compile(
    r"^(?:\(?[1-9]\d*\)|[1-9]\d*[.)]|First(?:ly)?|Second(?:ly)?|Third(?:ly)?|Finally)(?:\s+|$)",
    re.IGNORECASE,
)
CITATION_GROUP_RE = re.compile(r"\[(\d+(?:\s*[-,]\s*\d+)*)\]")
CITATION_CHAIN_RE = re.compile(r"(?:\[\d+(?:\s*[-,]\s*\d+)*\][,;]?\s*){4,}")
AUTHOR_CITATION_RE = re.compile(
    r"\b[A-Z][A-Za-z'’-]+(?:\s+(?:and|&|和)\s*[A-Z][A-Za-z'’-]+)?(?:\s+et\s+al\.|等)?\s*\[\d",
)
OBJECT_CITATION_RE = re.compile(
    r"(?:\b[A-Za-z][A-Za-z-]{2,}(?:\s+[A-Za-z][A-Za-z-]{2,}){0,5}|[\u3400-\u9fff]{2,12})\s*\[\d"
)
SENTENCE_END_CITATION_RE = re.compile(r"\[[0-9,\s-]+\]\s*[。！？；.!?:：;]?\s*$")
ACRONYM_DEFINITION_RE = re.compile(
    r"\b(?P<full>[A-Za-z][A-Za-z-]*(?:\s+[A-Za-z][A-Za-z-]*){1,10})\s*\((?P<acro>[A-Z][A-Z0-9-]{1,9})\)"
)
ACRONYM_RE = re.compile(r"\b[A-Z](?:[A-Z0-9-]{0,7}[A-Z0-9])\b")
ACRONYM_WHITELIST = {
    "AC", "DC", "AI", "CPU", "GPU", "PDF", "DOCX", "DOI", "IEEE", "CI", "ANOVA",
    "RMSE", "MAE", "MSE", "F1", "AUROC", "AUC", "USA", "UK", "SI", "API", "XML", "OR",
    "Q1", "Q3", "L1", "PASS", "FAIL", "READY", "EXACT", "TRUNCATED", "CSV", "JSON",
}
FIRST_PERSON_RE = re.compile(r"\b(?:we|our|ours|us)\b", re.IGNORECASE)

REPORT_STYLE_OPENING_RE = re.compile(
    r"^(?:As (?:shown|reported|listed|presented) in |(?:Figure|Fig\.|Table)\s+\d+\s+"
    r"(?:shows?|presents?|illustrates?|reports?)\b|From (?:Figure|Fig\.|Table)\s+\d+)",
    re.IGNORECASE,
)
AI_TEMPLATE_RE = re.compile(
    r"\b(?:there are (?:three|four|several) main (?:contributions|findings)|"
    r"the following conclusions can be drawn|in today's rapidly evolving|"
    r"plays a crucial role|in conclusion, this paper has|it is undeniable that)\b",
    re.IGNORECASE,
)
GENERIC_TRANSITION_RE = re.compile(
    r"^(?:First(?:ly)?|Second(?:ly)?|Third(?:ly)?|Finally|Moreover|Furthermore|Additionally|"
    r"In summary|In conclusion|Overall|Therefore)\b",
    re.IGNORECASE,
)
NUMBER_RE = re.compile(r"(?<![A-Za-z])[-+]?\d+(?:\.\d+)?(?:\s*(?:%|ms|s|m|km|J|kJ|W|Wh))?")
OPENING_TOKEN_RE = re.compile(r"^(?:The|This|These|Our|Results|Figure|Table|As|Compared|Under)\b", re.IGNORECASE)

SIGNIFICANCE_RE = re.compile(r"\bsignifican(?:t|tly)\b", re.IGNORECASE)
STATISTICAL_SUPPORT_RE = re.compile(
    r"\b(p\s*[<=>]|adjusted\s+p|confidence interval|\bCI\b|Holm|bootstrap|"
    r"permutation test|Wilcoxon|Mann[- ]Whitney|t[- ]test|ANOVA|effect size|Cohen|Hedges|Table\s+\d+)\b",
    re.IGNORECASE,
)

EVIDENCE_WORDS = re.compile(
    r"\b(Table|Figure|Fig\.|Equation|Dice|IoU|RMSE|accuracy|precision|recall|F1|AUROC|"
    r"mean|standard deviation|confidence interval|p\s*[<=>]|%|\d+\.\d+|\d+%)\b",
    re.IGNORECASE,
)
MECHANISM_WORDS = re.compile(
    r"\b(because|therefore|indicates|suggests|attributable|due to|whereas|although|"
    r"mechanism|module|component|ablation|diagnostic|limitation|bottleneck|explains|"
    r"reflects?|supports?|requires?|necessitating|confirms?|arises?|enables?|allows?|limits?|"
    r"reliability|variation|"
    r"sensitivity|consistent|inconsistent|boundary|constraint|trade-off)\b",
    re.IGNORECASE,
)
FIG_TAB_REF = re.compile(r"(?:\b(?:Figure|Fig\.|Table)\s+\d+\b|(?:图|表)\s*\d+)", re.IGNORECASE)
FIG_CAPTION_RE = re.compile(
    r"^(?:(?:Figure|Fig\.)\s+\d+|图\s*\d+)"
    r"(?:\s*[\.:：]\s*|\s+(?!(?:shows?|presents?|illustrates?|reports?|depicts?|indicates?|demonstrates?|"
    r"显示|表明|给出|展示)\b))",
    re.IGNORECASE,
)
TAB_CAPTION_RE = re.compile(
    r"^(?:Table\s+\d+|表\s*\d+)"
    r"(?:\s*[\.:：]\s*|\s+(?!(?:shows?|presents?|reports?|lists?|summarizes?|"
    r"显示|表明|给出|列出|汇总)\b))",
    re.IGNORECASE,
)


def detect_section(text: str, current: str) -> str:
    lowered = text.lower().strip().rstrip(":")
    bare = re.sub(r"^(?:[ivx]+|\d+(?:\.\d+)*)[.)]?\s+", "", lowered)
    chinese_headings = (
        (r"^摘\s*要", "abstract"),
        (r"^引言$", "introduction"),
        (r"^(?:相关工作|文献综述)$", "related work"),
        (r"^(?:主要贡献|本文贡献)$", "contributions"),
        (r"^(?:系统模型与所提方法|方法|方法学)$", "methodology"),
        (r"^(?:仿真设置与评价协议|仿真设置|实验设置)$", "experimental settings"),
        (r"^(?:仿真结果与讨论|实验结果与讨论|结果与讨论|仿真结果|实验结果)$", "results"),
        (r"^(?:结论|总结)$", "conclusion"),
        (r"^数据与代码可用性声明$", "data and code availability"),
        (r"^参考文献$", "references"),
        (r"^(?:项目概述与研究任务|项目概述|研究任务)$", "project overview"),
        (r"^(?:研究对象、资料与数据|研究对象与数据|资料与数据)$", "data"),
        (r"^(?:数学模型与求解|建模与求解)$", "methodology"),
        (r"^(?:分配结果与容量约束|工况变化与敏感性分析|技术结果与分析)$", "results"),
        (r"^(?:独立验证与技术评价|模型验证|独立验证)$", "validation"),
        (r"^(?:结论及工程建议|结论与建议|工程建议)$", "conclusion"),
        (r"^附录(?:\s*[A-ZＡ-Ｚ一二三四五六七八九十])?.*$", "appendix"),
    )
    for pattern, canonical in chinese_headings:
        if re.search(pattern, bare):
            return canonical
    section_names = (
        "abstract", "introduction", "background and motivation", "related works", "related work",
        "contributions", "methodology", "methods", "case study", "experiments", "experimental settings",
        "results", "discussion", "limitations", "conclusion", "conclusions", "data and code availability",
        "project overview", "data", "validation", "appendix", "appendices",
    )
    for name in section_names:
        if bare == name or bare.startswith(name + " "):
            return name
    return current


def reference_heading(text: str) -> bool:
    bare = re.sub(r"^(?:[ivx]+|\d+(?:\.\d+)*)[.)]?\s+", "", text.strip(), flags=re.I)
    return bare.lower().rstrip(":：") in {"references", "参考文献"}


def inline_abstract(text: str) -> bool:
    return bool(re.match(r"^(?:abstract|摘\s*要)\s*[:：]\s*\S", text, re.I))


def load_paragraphs(path: Path) -> list[dict]:
    """Retain section ancestry and structural breaks, including plain-style headings."""
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn
    doc = Document(path)
    paragraphs = []
    parent = section = "front_matter"
    boundary = False
    in_references = False
    index = 0
    for element in doc.element.body:
        if element.tag == qn("w:tbl"):
            boundary = True
            continue
        if element.tag != qn("w:p"):
            continue
        index += 1
        para = Paragraph(element, doc)
        text = " ".join(para.text.split())
        has_object = bool(element.xpath(".//w:object | .//w:drawing | .//m:oMath"))
        if not text:
            boundary = boundary or has_object
            continue
        style = para.style.name if para.style else ""
        if reference_heading(text):
            in_references = True
            boundary = True
            continue
        detected = detect_section(text, "")
        if inline_abstract(text):
            detected = "abstract"
        numbered = re.match(r"^(\d+(?:\.\d+)*)(?:\.)?\s+", text)
        heading = bool(style.lower().startswith("heading") or (
            numbered and len(text.split()) < 20 and not text.endswith((".", ";", "。", "；", "！", "？"))) or (
            detected and len(text.split()) < 9 and not text.endswith((".", ";", "。", "；", "！", "？"))))
        if inline_abstract(text):
            heading = False
            parent = section = "abstract"
        if in_references:
            if heading and detected in {"appendix", "appendices"}:
                in_references = False
                boundary = True
            else:
                continue
        if heading:
            depth = numbered.group(1).count(".") if numbered else 0
            if detected in {"abstract", "introduction", "methodology", "methods", "case study", "results", "experiments", "conclusion", "conclusions", "discussion", "limitations", "data and code availability", "project overview", "data", "validation", "appendix", "appendices"} and depth == 0:
                parent = "results" if detected == "case study" else detected
                section = parent
            elif detected in {"related work", "related works", "background and motivation", "contributions"}:
                parent = "introduction"
                section = parent + "/" + detected
            else:
                section = parent + "/" + text.lower()
        paragraphs.append({"index": index, "style": style, "section": section, "parent_section": parent,
                           "text": text, "is_heading": heading, "boundary_before": boundary or heading,
                           "contains_equation": has_object})
        boundary = has_object or bool(FIG_CAPTION_RE.match(text) or TAB_CAPTION_RE.match(text))
    return paragraphs


def find_terms(text: str, terms: list[str]) -> list[str]:
    return [
        term
        for term in terms
        if re.search(r"(?<!\w)" + re.escape(term) + r"(?!\w)", text, flags=re.IGNORECASE)
    ]


def load_editorial_blocks(path: Path, paragraphs: list[dict]) -> list[dict]:
    """Screen visible body/captions/table text without treating cells as prose transitions."""
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    from docx.oxml.ns import qn
    doc = Document(path)
    by_index = {p["index"]: p for p in paragraphs}
    blocks = []
    paragraph_index = table_index = 0
    section = "front_matter"
    in_references = False
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            paragraph_index += 1
            para = Paragraph(element, doc)
            text = " ".join(para.text.split())
            if reference_heading(text):
                in_references = True
                continue
            meta = by_index.get(paragraph_index, {})
            if in_references:
                if meta.get("is_heading") and meta.get("parent_section") in {"appendix", "appendices"}:
                    in_references = False
                else:
                    continue
            if not text:
                if element.xpath(".//w:object | .//w:drawing | .//m:oMath"):
                    blocks.append({"location": f"object:{paragraph_index}", "text": "", "section": section, "kind": "equation"})
                continue
            section = meta.get("section", section)
            kind = "heading" if meta.get("is_heading") else ("caption" if FIG_CAPTION_RE.match(text) or TAB_CAPTION_RE.match(text) else "prose")
            # A labeled note attached to an actual table/caption is supporting
            # text, not the preceding paragraph of a body argument. It still
            # receives all concrete editorial rules in editorial_screen.
            if (kind == "prose" and blocks
                    and blocks[-1].get("kind") in {"table", "caption", "note"}
                    and re.match(r"^(?:Notes?\s*[:.]|注\s*[:：])", text, re.I)):
                kind = "note"
            blocks.append({"location": f"paragraph:{paragraph_index}", "paragraph": paragraph_index,
                           "text": text, "section": section, "kind": kind,
                           "boundary": meta.get("boundary_before", False)})
        elif element.tag == qn("w:tbl"):
            table_index += 1
            if in_references:
                continue
            table = Table(element, doc)
            seen_cells = set()
            for row_index, row in enumerate(table.rows, 1):
                for cell_index, cell in enumerate(row.cells, 1):
                    if cell._tc in seen_cells:
                        continue
                    seen_cells.add(cell._tc)
                    text = " ".join(cell.text.split())
                    if text:
                        blocks.append({"location": f"table:{table_index}/row:{row_index}/cell:{cell_index}",
                                       "table": table_index, "row": row_index, "cell": cell_index,
                                       "text": text, "section": section, "kind": "table", "boundary": True})
    return blocks


def find_context_terms(text: str, terms: list[str]) -> list[str]:
    hits = []
    for term in terms:
        if re.search(r"[\u3400-\u9fff]", term):
            matched = term in text
        else:
            matched = bool(re.search(r"(?<!\w)" + re.escape(term) + r"(?!\w)", text, flags=re.IGNORECASE))
        if matched:
            hits.append(term)
    return hits


def normalized_words(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


def detect_language(paragraphs: list[dict]) -> str:
    text = "".join(p["text"] for p in paragraphs if not p.get("is_heading"))
    chinese = len(re.findall(r"[\u3400-\u9fff]", text))
    latin = len(re.findall(r"[A-Za-z]", text))
    return "chinese" if chinese >= max(20, latin // 2) else "english"


def prose_units(text: str, language: str) -> int:
    if language == "chinese":
        return len(re.findall(r"[\u3400-\u9fff]|[A-Za-z0-9]+", text))
    return len(normalized_words(text))


def chinese_sentence_skeleton(sentence: str) -> str:
    cleaned = re.sub(r"[-+]?\d+(?:\.\d+)?(?:%|[A-Za-z]+)?", "数值", sentence)
    cleaned = re.sub(r"(?:图|表)\s*\d+", "图表", cleaned)
    tokens = re.findall(r"[\u3400-\u9fff]{1,2}", cleaned)
    return "".join(tokens[:4]) if len(tokens) >= 4 else ""


def section_is(section: str, name: str) -> bool:
    return bool(re.search(rf"\b{re.escape(name)}s?\b", section, re.IGNORECASE))


def citation_group_size(group: str) -> int:
    total = 0
    for part in re.split(r"\s*,\s*", group):
        if "-" in part:
            try:
                start, end = [int(x.strip()) for x in part.split("-", 1)]
                total += max(1, end - start + 1)
            except ValueError:
                total += 1
        else:
            total += 1
    return total


def sentence_skeleton(sentence: str) -> str:
    cleaned = re.sub(r"\[[0-9,\s-]+\]", " citation ", sentence.lower())
    cleaned = re.sub(r"[-+]?\d+(?:\.\d+)?(?:%|\s*[a-z]+)?", " number ", cleaned)
    words = re.findall(r"[a-z]+", cleaned)
    return " ".join(words[:5]) if len(words) >= 5 else ""


def normalize_contribution(text: str) -> list[str]:
    text = re.sub(r"^\s*\(?\d+\)?[.)]?\s*", "", text.lower())
    return re.findall(r"[a-z0-9]+", text)


def split_sentences(text: str) -> list[str]:
    """Split English and Chinese prose without separating a trailing citation."""
    return [
        item.strip()
        for item in re.split(r"(?<=[。！？!?])\s*|(?<=[.;])\s+(?=[A-Z\u3400-\u9fff])", text)
        if item.strip()
    ]


def citation_modes_for_sentence(sentence: str) -> set[str]:
    """Return one deterministic dominant semantic attachment mode per sentence."""
    if AUTHOR_CITATION_RE.search(sentence):
        return {"author_attribution"}
    if SENTENCE_END_CITATION_RE.search(sentence):
        return {"sentence_end_attribution"}
    if OBJECT_CITATION_RE.search(sentence):
        return {"method_or_object_attribution"}
    if CITATION_GROUP_RE.search(sentence):
        return {"other_attribution"}
    return set()


def load_evidence_role_map(path: Path | None) -> dict:
    if path is None:
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


EXPERIMENT_AGENDA_RE = re.compile(
    r"this experiment is intended to|the purpose of this diagnostic is|this comparison asks whether|"
    r"this distinction is important because|was included to isolate|does not replace|"
    r"the diagnostic is intentionally evaluated|should not be interpreted as", re.I)


def section_diagnostics(paragraphs, path, manuscript_mode, role_map, language="auto"):
    """Advisory pattern checks; artifacts and headings break short-paragraph runs."""
    language = detect_language(paragraphs) if language == "auto" else language
    warnings = []
    sections = defaultdict(list)
    short_run = []
    previous_section = None

    def finish_run():
        if len(short_run) >= 3:
            warnings.append({"type": "paragraph_fragmentation", "paragraph": short_run[0]["index"],
                             "paragraphs": [p["index"] for p in short_run],
                             "note": "Adjacent short Results paragraphs; merge only when the scientific question is shared."})
        short_run.clear()

    for para in paragraphs:
        text = para["text"]
        result = section_is(para["section"], "result") or section_is(para["section"], "experiment")
        caption = bool(FIG_CAPTION_RE.match(text) or TAB_CAPTION_RE.match(text))
        if para.get("boundary_before") or para["section"] != previous_section or caption or not result:
            finish_run()
        previous_section = para["section"]
        if para.get("is_heading"):
            continue
        sections[para["section"]].append(para)
        if not result or caption:
            continue
        short_limit = 80 if language == "chinese" else 45
        if prose_units(text, language) <= short_limit and len(split_sentences(text)) <= 2 and not para.get("contains_equation"):
            short_run.append(para)
        else:
            finish_run()
        agenda = EXPERIMENT_AGENDA_RE.findall(text)
        if len(agenda) >= 2:
            warnings.append({"type": "experimental_report_prose", "paragraph": para["index"], "hits": agenda})
        refs = Counter(m.group().lower() for m in FIG_TAB_REF.finditer(text))
        repeated = {k: v for k, v in refs.items() if v >= 3}
        if repeated:
            warnings.append({"type": "repeated_figure_reference", "paragraph": para["index"], "references": repeated})
        if re.search(r"correlat|associat", text, re.I) and re.search(r"proves? .{0,45}caus|causally establishes?|causes?", text, re.I) and not role_map.get("causal_identification"):
            warnings.append({"type": "unsupported_causal_language", "paragraph": para["index"], "text": text})
    finish_run()
    for section, items in sections.items():
        if not (section_is(section, "result") or section_is(section, "experiment")):
            continue
        prose = [p for p in items if not (FIG_CAPTION_RE.match(p["text"]) or TAB_CAPTION_RE.match(p["text"]))]
        refs = {m.group().lower() for p in items for m in FIG_TAB_REF.finditer(p["text"])}
        words = sum(prose_units(p["text"], language) for p in prose)
        unit_floor = 70 if language == "chinese" else 35
        if len(refs) >= 3 and len(prose) <= 2 and words < unit_floor * len(refs):
            warnings.append({"type": "analysis_undercompression", "paragraph": items[0]["index"],
                             "section": section, "artifacts": sorted(refs), "prose_words": words,
                             "note": "Figure/table-rich section has little analysis; restore reasoning, not numerical replay."})
        agenda_count = sum(len(EXPERIMENT_AGENDA_RE.findall(p["text"])) for p in prose)
        if agenda_count >= 3 and not any(w["type"]=="experimental_report_prose" and w["paragraph"] in {p["index"] for p in prose} for w in warnings):
            warnings.append({"type": "experimental_report_prose", "paragraph": prose[0]["index"], "section": section, "hits": agenda_count})

    # Compare actual table values, rather than equating every numerical paragraph with replay.
    numeric_rows = []
    for table in Document(path).tables:
        values = {n for row in table.rows for cell in row.cells for n in NUMBER_RE.findall(cell.text)}
        if len(values) >= 6: numeric_rows.append(values)
    for para in paragraphs:
        if not section_is(para["section"], "result") or para.get("is_heading"):
            continue
        vals = set(NUMBER_RE.findall(para["text"]))
        if len(vals) >= 6 and any(len(vals & row) >= 6 and len(vals & row)/len(vals) >= .75 for row in numeric_rows):
            warnings.append({"type": "table_replay", "paragraph": para["index"], "note": "Many prose values repeat a table; retain only decisive quantitative anchors."})

    seen = {}
    for para in paragraphs:
        if not any(section_is(para["section"], s) for s in ("methodology", "methods")):
            continue
        for m in re.finditer(r"(?:where|and)\s+\$?([A-Za-z](?:_\{?\w+\}?)?)\$?\s+(?:denotes|is defined as|represents)\s+([^.;]+)", para["text"], re.I):
            key = (para["section"], m.group(1), " ".join(m.group(2).lower().split()))
            if key in seen:
                warnings.append({"type": "stable_notation_redefinition", "paragraph": para["index"], "first_paragraph": seen[key], "symbol": m.group(1)})
            else: seen[key] = para["index"]
    for para in paragraphs:
        if "contribution" in para["section"] and not para.get("is_heading"):
            if re.match(r"\s*(?:\(?\d+[.)]\s*)?(?:we |this work |the study )?(?:conduct|evaluat|perform|test|validat|achiev)", para["text"], re.I) and re.search(r"experiment|dataset|baseline|significance|accuracy", para["text"], re.I):
                warnings.append({"type": "experiments_as_contribution", "paragraph": para["index"], "text": para["text"]})
    # A conclusion synthesizes established science. Missing-experiment caveats
    # and future-work padding are not mandatory Journal paragraph content.
    return warnings


def inspect_writing(
    path: Path,
    manuscript_mode: str = "project",
    evidence_role_map: dict | None = None,
    acronym_whitelist: set[str] | None = None,
    language: str = "auto",
    review_decisions: dict | None = None,
) -> dict:
    paragraphs = load_paragraphs(path)
    language = detect_language(paragraphs) if language == "auto" else language
    warnings: list[dict] = []
    counts = {
        "paragraphs_checked": len(paragraphs),
        "prompt_residue_hits": 0,
        "overclaim_hits": 0,
        "weak_phrase_hits": 0,
        "unsupported_significance_hits": 0,
        "long_paragraphs": 0,
        "figure_table_reference_paragraphs": 0,
        "shallow_figure_table_reference_paragraphs": 0,
        "defensive_phrase_hits": 0,
        "report_style_openings": 0,
        "report_style_density_flags": 0,
        "repeated_opening_runs": 0,
        "number_dumping_paragraphs": 0,
        "ai_template_hits": 0,
        "generic_transition_paragraphs": 0,
        "generic_transition_density_flags": 0,
        "weak_mechanism_paragraphs": 0,
        "conclusion_duplication_pairs": 0,
        "report_style_list_paragraphs": 0,
        "contrast_phrase_hits": 0,
        "contrast_density_flags": 0,
        "mechanical_enumeration_hits": 0,
        "mechanical_enumeration_density_flags": 0,
        "defensive_density_flags": 0,
        "citation_stack_hits": 0,
        "undefined_acronym_hits": 0,
        "repeated_acronym_definition_hits": 0,
        "duplicate_contribution_pairs": 0,
        "repetitive_sentence_skeletons": 0,
        "multi_panel_interpretation_gaps": 0,
        "evidence_role_mismatches": 0,
        "journal_first_person_hits": 0,
        "experiment_governance_residue_hits": 0,
        "project_lineage_residue_hits": 0,
        "citation_attribution_monotony_flags": 0,
        "contribution_external_citation_hits": 0,
        "invented_quote_concept_hits": 0,
        "em_dash_chain_hits": 0,
        "abstract_noun_stack_hits": 0,
        "hollow_report_prose_hits": 0,
        "process_language_leakage_hits": 0,
        "definite_placeholder_hits": 0,
        "chinese_repetitive_sentence_skeletons": 0,
    }

    citation_modes = Counter()
    introduction_citation_modes = Counter()
    introduction_citation_bearing_sentence_count = 0
    introduction_author_sentence_count = 0
    skeleton_locations: defaultdict[str, list[int]] = defaultdict(list)
    chinese_skeleton_locations: defaultdict[str, list[int]] = defaultdict(list)
    definitions_by_domain: defaultdict[str, defaultdict[str, list[tuple[int, str]]]] = defaultdict(lambda: defaultdict(list))
    acronyms_by_domain: defaultdict[str, defaultdict[str, list[int]]] = defaultdict(lambda: defaultdict(list))
    contrast_paragraphs = 0
    enumerated_paragraphs: list[dict] = []
    allowed_acronyms = ACRONYM_WHITELIST | {item.upper() for item in (acronym_whitelist or set())}

    opening_run: list[dict] = []
    last_opening = ""

    for para in paragraphs:
        text = para["text"]
        word_count = prose_units(text, language)
        is_caption = bool(FIG_CAPTION_RE.match(text) or TAB_CAPTION_RE.match(text))
        prompt_hits = find_terms(text, PROMPT_RESIDUE)
        overclaim_hits = find_terms(text, OVERCLAIM_TERMS)
        weak_hits = find_terms(text, WEAK_PHRASES)
        defensive_hits = find_terms(text, DEFENSIVE_PHRASES)
        governance_hits = []
        if para.get("parent_section", para["section"]) in GOVERNANCE_SENSITIVE_SECTIONS:
            governance_hits = find_context_terms(text, EXPERIMENT_GOVERNANCE_RESIDUE)
        lineage_hits = find_context_terms(text, PROJECT_LINEAGE_RESIDUE)
        invented_quotes = INVENTED_QUOTE_RE.findall(text) if language == "chinese" else []
        em_dash_chain = language == "chinese" and text.count("—") >= 2
        abstract_stack = CHINESE_ABSTRACT_STACK_RE.search(text) if language == "chinese" else None
        hollow_report = CHINESE_HOLLOW_REPORT_RE.search(text) if language == "chinese" else None
        process_leakage = CHINESE_PROCESS_LEAKAGE_RE.search(text) if language == "chinese" else None
        placeholders = DEFINITE_PLACEHOLDER_RE.findall(text)
        if prompt_hits:
            counts["prompt_residue_hits"] += len(prompt_hits)
            warnings.append({"paragraph": para["index"], "type": "prompt_residue", "hits": prompt_hits, "text": text})
        if overclaim_hits:
            counts["overclaim_hits"] += len(overclaim_hits)
            warnings.append({"paragraph": para["index"], "type": "overclaim_risk", "hits": overclaim_hits, "text": text})
        if weak_hits:
            counts["weak_phrase_hits"] += len(weak_hits)
            warnings.append({"paragraph": para["index"], "type": "weak_phrase", "hits": weak_hits, "text": text})
        if defensive_hits:
            counts["defensive_phrase_hits"] += len(defensive_hits)
            warnings.append({"paragraph": para["index"], "type": "defensive_prose", "hits": defensive_hits, "text": text})
        if governance_hits:
            counts["experiment_governance_residue_hits"] += len(governance_hits)
            warnings.append(
                {
                    "paragraph": para["index"],
                    "type": "experiment_governance_residue",
                    "section": para["section"],
                    "hits": governance_hits,
                    "text": text,
                }
            )
        if lineage_hits:
            counts["project_lineage_residue_hits"] += len(lineage_hits)
            warnings.append(
                {
                    "paragraph": para["index"],
                    "type": "project_lineage_residue",
                    "section": para["section"],
                    "hits": lineage_hits,
                    "text": text,
                }
            )
        if invented_quotes:
            counts["invented_quote_concept_hits"] += len(invented_quotes)
            warnings.append({"paragraph": para["index"], "type": "invented_quote_concept", "hits": invented_quotes, "text": text})
        if em_dash_chain:
            counts["em_dash_chain_hits"] += 1
            warnings.append({"paragraph": para["index"], "type": "em_dash_chain", "text": text})
        if abstract_stack:
            counts["abstract_noun_stack_hits"] += 1
            warnings.append({"paragraph": para["index"], "type": "abstract_noun_stack", "text": text})
        if hollow_report and not (NUMBER_RE.search(text) or re.search(r"由于|因为|导致|约束|数据|实测|计算", text)):
            counts["hollow_report_prose_hits"] += 1
            warnings.append({"paragraph": para["index"], "type": "hollow_report_prose", "text": text})
        if process_leakage:
            counts["process_language_leakage_hits"] += 1
            warnings.append({"paragraph": para["index"], "type": "process_language_leakage", "text": text})
        if placeholders:
            counts["definite_placeholder_hits"] += len(placeholders)
            warnings.append({"paragraph": para["index"], "type": "definite_placeholder", "hits": placeholders, "text": text, "severity": "blocker"})
        contrast_hits = CONTRAST_PHRASE_RE.findall(text)
        if contrast_hits:
            counts["contrast_phrase_hits"] += len(contrast_hits)
            contrast_paragraphs += 1
        if not is_caption and not para.get("is_heading", False) and MECHANICAL_ENUMERATION_RE.search(text):
            counts["mechanical_enumeration_hits"] += 1
            enumerated_paragraphs.append(para)
        if manuscript_mode == "journal" and not is_caption:
            first_person_hits = FIRST_PERSON_RE.findall(text)
            if first_person_hits:
                counts["journal_first_person_hits"] += len(first_person_hits)
                warnings.append({"paragraph": para["index"], "type": "journal_first_person", "hits": first_person_hits, "text": text})
        stack_found = False
        for group in CITATION_GROUP_RE.findall(text):
            size = citation_group_size(group)
            if size >= 4:
                stack_found = True
                counts["citation_stack_hits"] += 1
                warnings.append({"paragraph": para["index"], "type": "citation_stack", "reference_count": size, "text": text})
        if not stack_found and CITATION_CHAIN_RE.search(text):
            size = len(CITATION_GROUP_RE.findall(CITATION_CHAIN_RE.search(text).group(0)))
            counts["citation_stack_hits"] += 1
            warnings.append({"paragraph": para["index"], "type": "citation_stack", "reference_count": size, "text": text})
        for sentence in split_sentences(text):
            sentence_modes = citation_modes_for_sentence(sentence)
            for mode in sentence_modes:
                citation_modes[mode] += 1
            if section_is(para["section"], "introduction") and not section_is(para["section"], "contribution") and CITATION_GROUP_RE.search(sentence):
                introduction_citation_bearing_sentence_count += 1
                if "author_attribution" in sentence_modes:
                    introduction_author_sentence_count += 1
                for mode in sentence_modes:
                    introduction_citation_modes[mode] += 1
        contribution_like = section_is(para["section"], "contribution") or bool(
            re.search(r"(?:本文|本研究)(?:的)?主要贡献(?:包括|如下)|our (?:main )?contributions?", text, re.IGNORECASE)
        )
        if contribution_like and CITATION_GROUP_RE.search(text):
            citation_count = len(CITATION_GROUP_RE.findall(text))
            counts["contribution_external_citation_hits"] += citation_count
            warnings.append(
                {
                    "paragraph": para["index"],
                    "type": "contribution_external_citation",
                    "citation_group_count": citation_count,
                    "text": text,
                }
            )

        domain = "abstract" if section_is(para["section"], "abstract") else "main_text"
        for match in ACRONYM_DEFINITION_RE.finditer(text):
            definitions_by_domain[domain][match.group("acro")].append((para["index"], match.group("full")))
        for acronym in ACRONYM_RE.findall(text):
            if acronym not in allowed_acronyms:
                acronyms_by_domain[domain][acronym].append(para["index"])
        for sentence in split_sentences(text):
            skeleton = sentence_skeleton(sentence)
            if skeleton:
                skeleton_locations[skeleton].append(para["index"])
            if language == "chinese":
                chinese_skeleton = chinese_sentence_skeleton(sentence)
                if chinese_skeleton:
                    chinese_skeleton_locations[chinese_skeleton].append(para["index"])
        if not is_caption and REPORT_STYLE_OPENING_RE.search(text):
            counts["report_style_openings"] += 1
            warnings.append({"paragraph": para["index"], "type": "report_style_opening", "text": text})
        if not is_caption and AI_TEMPLATE_RE.search(text):
            counts["ai_template_hits"] += 1
            warnings.append({"paragraph": para["index"], "type": "ai_template_phrase", "text": text})
        if not is_caption and GENERIC_TRANSITION_RE.search(text):
            counts["generic_transition_paragraphs"] += 1
        if (
            not is_caption
            and (section_is(para["section"], "result") or section_is(para["section"], "discussion"))
            and para["style"].lower().startswith("list")
        ):
            counts["report_style_list_paragraphs"] += 1
            warnings.append(
                {"paragraph": para["index"], "type": "report_style_list", "style": para["style"], "text": text}
            )

        numbers = NUMBER_RE.findall(text)
        protected_section = "abstract" in para["section"] or "conclusion" in para["section"]
        if len(numbers) >= 8 and not protected_section and not MECHANISM_WORDS.search(text) and not STATISTICAL_SUPPORT_RE.search(text):
            counts["number_dumping_paragraphs"] += 1
            warnings.append(
                {"paragraph": para["index"], "type": "number_dumping", "number_count": len(numbers), "text": text}
            )

        opening_match = None if is_caption else OPENING_TOKEN_RE.match(text)
        opening = opening_match.group(0).lower() if opening_match else ""
        if opening and opening == last_opening:
            opening_run.append(para)
        else:
            if len(opening_run) >= 3:
                counts["repeated_opening_runs"] += 1
                warnings.append(
                    {
                        "paragraph": opening_run[0]["index"],
                        "type": "repeated_paragraph_opening",
                        "opening": last_opening,
                        "paragraphs": [item["index"] for item in opening_run],
                    }
                )
            opening_run = [para] if opening else []
        last_opening = opening
        if SIGNIFICANCE_RE.search(text) and not STATISTICAL_SUPPORT_RE.search(text):
            counts["unsupported_significance_hits"] += 1
            warnings.append(
                {
                    "paragraph": para["index"],
                    "type": "unsupported_significance_language",
                    "text": text,
                }
            )
        long_limit = 420 if language == "chinese" else 240
        if word_count > long_limit and not is_caption and not MECHANISM_WORDS.search(text):
            counts["long_paragraphs"] += 1
            warnings.append({"paragraph": para["index"], "type": "long_paragraph", "word_count": word_count, "text": text[:500]})

        if is_caption:
            continue
        if FIG_TAB_REF.search(text):
            counts["figure_table_reference_paragraphs"] += 1
            has_evidence = bool(NUMBER_RE.search(text) or EVIDENCE_WORDS.search(text))
            has_mechanism = bool(MECHANISM_WORDS.search(text))
            artifact_led = bool(REPORT_STYLE_OPENING_RE.search(text))
            vague_result = bool(re.search(r"\b(?:shows?|presents?|reports?)\s+(?:the\s+)?results?\b", text, re.IGNORECASE))
            very_short = word_count <= 12
            if (artifact_led and not (has_evidence and has_mechanism)) or vague_result or very_short:
                counts["shallow_figure_table_reference_paragraphs"] += 1
                warnings.append(
                    {
                        "paragraph": para["index"],
                        "type": "shallow_figure_table_discussion",
                        "has_evidence": has_evidence,
                        "has_mechanism": has_mechanism,
                        "text": text,
                    }
                )

    if len(opening_run) >= 3:
        counts["repeated_opening_runs"] += 1
        warnings.append(
            {
                "paragraph": opening_run[0]["index"],
                "type": "repeated_paragraph_opening",
                "opening": last_opening,
                "paragraphs": [item["index"] for item in opening_run],
            }
        )

    prose_paragraphs = [
        para for para in paragraphs
        if not (FIG_CAPTION_RE.match(para["text"]) or TAB_CAPTION_RE.match(para["text"]))
        and not para.get("is_heading", False)
        and prose_units(para["text"], language) >= (16 if language == "chinese" else 8)
    ]
    transition_density = counts["generic_transition_paragraphs"] / max(1, len(prose_paragraphs))
    if counts["generic_transition_paragraphs"] >= 3 and transition_density >= 0.15:
        counts["generic_transition_density_flags"] = 1
        warnings.append(
            {
                "paragraph": 0,
                "type": "generic_transition_density",
                "paragraph_count": counts["generic_transition_paragraphs"],
                "density": round(transition_density, 3),
            }
        )

    report_style_density = counts["report_style_openings"] / max(1, len(prose_paragraphs))
    if counts["report_style_openings"] >= 3 and report_style_density >= 0.12:
        counts["report_style_density_flags"] = 1
        warnings.append({"paragraph": 0, "type": "report_style_density", "paragraph_count": counts["report_style_openings"], "density": round(report_style_density, 3)})

    contrast_density = contrast_paragraphs / max(1, len(prose_paragraphs))
    if counts["contrast_phrase_hits"] >= 3 and contrast_density >= 0.10:
        counts["contrast_density_flags"] = 1
        warnings.append({"paragraph": 0, "type": "formulaic_contrast_density", "hit_count": counts["contrast_phrase_hits"], "density": round(contrast_density, 3)})

    non_contribution_enumeration = [p for p in enumerated_paragraphs if "contribution" not in p["section"]]
    enumeration_density = len(non_contribution_enumeration) / max(1, len(prose_paragraphs))
    if len(non_contribution_enumeration) >= 3 and enumeration_density >= 0.08:
        counts["mechanical_enumeration_density_flags"] = 1
        warnings.append({"paragraph": non_contribution_enumeration[0]["index"], "type": "mechanical_enumeration_density", "paragraphs": [p["index"] for p in non_contribution_enumeration], "density": round(enumeration_density, 3)})

    defensive_density = counts["defensive_phrase_hits"] / max(1, len(prose_paragraphs))
    defensive_outside_scope_sections = sum(
        1 for p in prose_paragraphs
        if find_terms(p["text"], DEFENSIVE_PHRASES)
        and not any(section_is(p["section"], s) for s in ("discussion", "limitation", "conclusion"))
    )
    if counts["defensive_phrase_hits"] >= 3 and (defensive_density >= 0.05 or defensive_outside_scope_sections >= 2):
        counts["defensive_density_flags"] = 1
        warnings.append({"paragraph": 0, "type": "defensive_prose_density", "hit_count": counts["defensive_phrase_hits"], "outside_scope_sections": defensive_outside_scope_sections, "density": round(defensive_density, 3)})

    introduction_author_ratio = introduction_author_sentence_count / max(
        1, introduction_citation_bearing_sentence_count
    )
    other_meaningful_modes = (
        introduction_citation_modes.get("method_or_object_attribution", 0)
        + introduction_citation_modes.get("sentence_end_attribution", 0)
        + introduction_citation_modes.get("other_attribution", 0)
    )
    if (
        introduction_citation_bearing_sentence_count >= 6
        and introduction_author_ratio > 0.70
        and other_meaningful_modes < 2
    ):
        counts["citation_attribution_monotony_flags"] = 1
        warnings.append(
            {
                "paragraph": 0,
                "type": "citation_attribution_monotony",
                "section": "introduction",
                "citation_bearing_sentence_count": introduction_citation_bearing_sentence_count,
                "author_attribution_ratio": round(introduction_author_ratio, 4),
                "other_meaningful_mode_count": other_meaningful_modes,
            }
        )

    for skeleton, locations in skeleton_locations.items():
        unique_locations = sorted(set(locations))
        if len(unique_locations) >= 3:
            counts["repetitive_sentence_skeletons"] += 1
            warnings.append({"paragraph": unique_locations[0], "type": "repetitive_sentence_skeleton", "skeleton": skeleton, "paragraphs": unique_locations})
    for skeleton, locations in chinese_skeleton_locations.items():
        unique_locations = sorted(set(locations))
        if len(unique_locations) >= 3:
            counts["chinese_repetitive_sentence_skeletons"] += 1
            warnings.append({"paragraph": unique_locations[0], "type": "chinese_repetitive_sentence_skeleton", "skeleton": skeleton, "paragraphs": unique_locations})

    for domain, acronyms in acronyms_by_domain.items():
        for acronym, locations in acronyms.items():
            if acronym not in definitions_by_domain[domain]:
                counts["undefined_acronym_hits"] += 1
                warnings.append({"paragraph": locations[0], "type": "undefined_acronym", "domain": domain, "acronym": acronym})
    for domain, definitions in definitions_by_domain.items():
        for acronym, items in definitions.items():
            if len(items) > 1:
                counts["repeated_acronym_definition_hits"] += len(items) - 1
                warnings.append({"paragraph": items[1][0], "type": "repeated_acronym_definition", "domain": domain, "acronym": acronym, "paragraphs": [x[0] for x in items]})

    contributions = [p for p in prose_paragraphs if "contribution" in p["section"] and re.match(r"^\s*\(?\d+\)?[.)]", p["text"])]
    for i, left in enumerate(contributions):
        left_words = normalize_contribution(left["text"])
        for right in contributions[i + 1:]:
            right_words = normalize_contribution(right["text"])
            if len(left_words) < 12 or len(right_words) < 12:
                continue
            jaccard = len(set(left_words) & set(right_words)) / max(1, len(set(left_words) | set(right_words)))
            sequence = SequenceMatcher(None, left_words, right_words, autojunk=False).ratio()
            if jaccard >= 0.72 or sequence >= 0.82:
                counts["duplicate_contribution_pairs"] += 1
                warnings.append({"paragraph": right["index"], "type": "duplicate_contribution", "other_paragraph": left["index"], "jaccard": round(jaccard, 3), "sequence_similarity": round(sequence, 3)})

    captions = [p for p in paragraphs if FIG_CAPTION_RE.match(p["text"])]
    body_text = "\n".join(p["text"] for p in prose_paragraphs if not FIG_CAPTION_RE.match(p["text"]))
    for caption in captions:
        number_match = re.match(r"^(?:Figure|Fig\.)\s+(\d+)", caption["text"], re.IGNORECASE)
        panels = sorted(set(re.findall(r"\(([a-z])\)", caption["text"], re.IGNORECASE)))
        if not number_match or len(panels) < 2:
            continue
        number = number_match.group(1)
        missing = [panel for panel in panels if not re.search(rf"(?:Figure|Fig\.)\s*{number}\s*\({panel}\)", body_text, re.IGNORECASE)]
        if missing:
            counts["multi_panel_interpretation_gaps"] += 1
            warnings.append({"paragraph": caption["index"], "type": "multi_panel_interpretation_gap", "figure": number, "missing_panels": missing})

    role_map = evidence_role_map or {}
    incompatible = {
        ("diagnostic", "primary_performance"), ("diagnostic", "generalization"),
        ("auxiliary", "primary_performance"), ("upper_bound", "deployable_performance"),
        ("qualitative", "quantitative_superiority"),
    }
    for artifact in role_map.get("artifacts", []):
        pair = (str(artifact.get("role", "")).lower(), str(artifact.get("claimed_as", "")).lower())
        if pair in incompatible:
            counts["evidence_role_mismatches"] += 1
            warnings.append({"paragraph": 0, "type": "evidence_role_mismatch", "artifact_id": artifact.get("id", ""), "role": pair[0], "claimed_as": pair[1]})

    result_paragraphs = [
        para for para in prose_paragraphs
        if section_is(para["section"], "result") and para["text"].lower() != "results"
    ]
    result_positions = {para["index"]: position for position, para in enumerate(result_paragraphs)}
    for para in result_paragraphs:
        text = para["text"]
        if len(normalized_words(text)) < 18 or not (NUMBER_RE.search(text) or EVIDENCE_WORDS.search(text)):
            continue
        position = result_positions[para["index"]]
        neighborhood = result_paragraphs[max(0, position - 1): position + 2]
        if not any(MECHANISM_WORDS.search(item["text"]) for item in neighborhood):
            counts["weak_mechanism_paragraphs"] += 1
            warnings.append(
                {"paragraph": para["index"], "type": "weak_mechanism_analysis", "text": text}
            )

    abstract_paragraphs = [
        para for para in prose_paragraphs
        if section_is(para["section"], "abstract") and para["text"].lower() != "abstract"
    ]
    conclusion_paragraphs = [
        para for para in prose_paragraphs
        if section_is(para["section"], "conclusion") and para["text"].lower() not in {"conclusion", "conclusions"}
    ]
    for conclusion in conclusion_paragraphs:
        conclusion_words = normalized_words(conclusion["text"])
        if len(conclusion_words) < 20:
            continue
        for abstract in abstract_paragraphs:
            abstract_words = normalized_words(abstract["text"])
            if len(abstract_words) < 20:
                continue
            union = set(conclusion_words) | set(abstract_words)
            jaccard = len(set(conclusion_words) & set(abstract_words)) / max(1, len(union))
            sequence = SequenceMatcher(None, conclusion_words, abstract_words, autojunk=False).ratio()
            if jaccard >= 0.65 or sequence >= 0.72:
                counts["conclusion_duplication_pairs"] += 1
                warnings.append(
                    {
                        "paragraph": conclusion["index"],
                        "type": "conclusion_duplicates_abstract",
                        "abstract_paragraph": abstract["index"],
                        "jaccard": round(jaccard, 3),
                        "sequence_similarity": round(sequence, 3),
                    }
                )

    new_warnings = section_diagnostics(paragraphs, path, manuscript_mode, role_map, language)
    warnings.extend(new_warnings)
    for warning in new_warnings:
        key = warning["type"] + "_flags"
        counts[key] = counts.get(key, 0) + 1

    score = 100
    score -= min(30, counts["prompt_residue_hits"] * 10)
    score -= min(25, counts["overclaim_hits"] * 5)
    score -= min(20, counts["weak_phrase_hits"] * 3)
    score -= min(20, counts["unsupported_significance_hits"] * 4)
    score -= min(15, counts["shallow_figure_table_reference_paragraphs"] * 2)
    score -= min(10, counts["long_paragraphs"] * 2)
    score -= min(10, counts["defensive_density_flags"] * 5)
    score -= min(10, counts["report_style_density_flags"] * 5)
    score -= min(10, counts["repeated_opening_runs"] * 2)
    score -= min(12, counts["number_dumping_paragraphs"] * 3)
    score -= min(10, counts["ai_template_hits"] * 3)
    score -= min(8, counts["generic_transition_density_flags"] * 5)
    score -= min(12, counts["weak_mechanism_paragraphs"] * 2)
    score -= min(12, counts["conclusion_duplication_pairs"] * 4)
    score -= min(10, counts["report_style_list_paragraphs"] * 2)
    score -= min(8, counts["contrast_density_flags"] * 4)
    score -= min(8, counts["mechanical_enumeration_density_flags"] * 4)
    score -= min(10, counts["citation_stack_hits"] * 2)
    score -= min(10, counts["undefined_acronym_hits"])
    score -= min(8, counts["repeated_acronym_definition_hits"] * 2)
    score -= min(15, counts["duplicate_contribution_pairs"] * 5)
    score -= min(8, counts["repetitive_sentence_skeletons"] * 2)
    score -= min(8, counts["multi_panel_interpretation_gaps"] * 2)
    score -= min(15, counts["evidence_role_mismatches"] * 5)
    score -= min(10, counts["journal_first_person_hits"])
    score -= min(30, counts["experiment_governance_residue_hits"] * 5)
    score -= min(20, counts["project_lineage_residue_hits"] * 5)
    score -= min(8, counts["citation_attribution_monotony_flags"] * 4)
    score -= min(10, counts["contribution_external_citation_hits"] * 2)
    score -= min(8, counts["invented_quote_concept_hits"] * 2)
    score -= min(8, counts["em_dash_chain_hits"] * 2)
    score -= min(8, counts["abstract_noun_stack_hits"] * 2)
    score -= min(10, counts["hollow_report_prose_hits"] * 3)
    score -= min(15, counts["process_language_leakage_hits"] * 5)
    score -= min(30, counts["definite_placeholder_hits"] * 10)
    score -= min(8, counts["chinese_repetitive_sentence_skeletons"] * 2)
    score = max(0, score)

    guidance = {
        "invented_quote_concept": ("引号包装了新造的抽象概念，定义与证据关系不清。", "去掉装饰性引号；若确为术语，给出可检验定义并在后文稳定使用。"),
        "em_dash_chain": ("同一段使用多组长破折号，可能用标点替代逻辑连接。", "按因果、条件或结果关系拆分句子；公式减号、数值范围和普通连字符无需修改。"),
        "abstract_noun_stack": ("多个抽象名词并列但缺少对象、动作或证据。", "保留真正需要的概念，并说明哪个对象发生了什么可观察变化。"),
        "hollow_report_prose": ("结论性评价缺少当前段落可追踪的计算、数据或机制。", "改写为具体对象、工况、量值或约束关系；没有证据时删除评价。"),
        "process_language_leakage": ("写作或修改过程进入正式报告正文。", "删除生成、提示或修改过程，直接陈述研究对象、方法和结果。"),
        "definite_placeholder": ("存在明确的未完成占位符。", "补全事实或从正式交付稿删除该占位。"),
        "chinese_repetitive_sentence_skeleton": ("相邻或多处句子使用相同开头和结构，削弱信息层次。", "依据事实关系改变主语和句法，只在结构确实平行时保留重复。"),
    }
    defaults = ("该模式需要结合相邻段落和证据上下文复核。", "在不改变科学含义的前提下，补充具体对象、条件和证据，或删除无信息表达。")
    paragraph_text = {p["index"]: p["text"] for p in paragraphs}
    for warning in warnings:
        reason, suggestion = guidance.get(warning["type"], defaults)
        warning.setdefault("original_text", warning.get("text") or paragraph_text.get(warning.get("paragraph"), ""))
        warning.setdefault("reason", reason)
        warning.setdefault("suggestion", suggestion)
        warning.setdefault("severity", "advisory")

    editorial = editorial_screen(load_editorial_blocks(path, paragraphs), manuscript_mode, review_decisions)
    return {
        "docx": str(path),
        "writing_quality_screen_score": score,
        "counts": counts,
        "citation_attribution_modes": dict(citation_modes),
        "introduction_citation_modes": dict(introduction_citation_modes),
        "introduction_citation_bearing_sentence_count": introduction_citation_bearing_sentence_count,
        "introduction_author_attribution_ratio": round(introduction_author_ratio, 4),
        "manuscript_mode": manuscript_mode,
        "language": language,
        "blocking_warning_count": sum(1 for warning in warnings if warning.get("severity") == "blocker"),
        "warnings": warnings,
        **editorial,
        "note": "Legacy warnings are draft diagnostics, not AI-authorship evidence. Journal/Conference/Project --require-clean rejects unresolved source-bound editorial candidates and invalid reviews. The ordinary draft command remains advisory; explicit placeholders remain blockers at final delivery.",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Screen manuscript prose for writing-quality risks.")
    parser.add_argument("docx_path")
    parser.add_argument("--json-out", default="")
    parser.add_argument("--manuscript-mode", choices=["journal", "conference", "project"], default="project")
    parser.add_argument("--language", choices=["auto", "english", "chinese"], default="auto")
    parser.add_argument("--evidence-role-map", default="", help="Optional JSON with artifacts[].id/role/claimed_as metadata.")
    parser.add_argument("--allow-acronym", action="append", default=[], help="Venue/domain-standard acronym to exclude from undefined-acronym diagnostics. Can be repeated.")
    parser.add_argument("--require-clean", action="store_true", help="Journal/Conference/Project final gate: reject unresolved editorial candidates or invalid review decisions. The ordinary draft command remains advisory.")
    parser.add_argument("--review-decisions", default="", help="JSON reviews with candidate_id, decision=scientific_content and a specific scientific reason; rewritten text must be rescanned.")
    args = parser.parse_args()
    role_map = load_evidence_role_map(Path(args.evidence_role_map).resolve()) if args.evidence_role_map else None
    result = inspect_writing(
        Path(args.docx_path).resolve(),
        manuscript_mode=args.manuscript_mode,
        evidence_role_map=role_map,
        acronym_whitelist=set(args.allow_acronym),
        language=args.language,
        review_decisions=read_reviews(args.review_decisions),
    )
    payload = json.dumps(result, indent=2, ensure_ascii=False)
    sys.stdout.buffer.write((payload + "\n").encode("utf-8", errors="replace"))
    if args.json_out:
        Path(args.json_out).write_text(payload, encoding="utf-8")
    if args.require_clean:
        return 2 if result["blocking_warning_count"] or result["unresolved_editorial_count"] or result["editorial_review_errors"] else 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
