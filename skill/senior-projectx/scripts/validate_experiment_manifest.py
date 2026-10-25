#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import sys
import uuid
import zipfile
from pathlib import Path
from typing import Any


VALID_STATUS = {"R&D draft", "Evidence candidate", "Submission candidate"}
VALID_TASK_STRUCTURES = {"independent_samples", "time_series", "spatiotemporal", "sequential_decision", "other"}
VALID_METHOD_STRUCTURES = {"single", "composite", "pipeline"}
VALID_GENERALIZATION_LEVELS = {
    "not_applicable", "within_site", "multi_site_in_distribution", "leave_site_out", "cross_region", "ood_regime"
}
VALID_CLAIM_TYPES = {
    "performance", "model_comparison", "stability", "generalization", "ood_generalization",
    "predictive_uncertainty", "composite_synergy", "deployment",
}
VALID_EVIDENCE_ROLES = {"raw_result", "analysis_script", "statistical_output", "table_or_figure", "model_checkpoint", "protocol"}
ROLE_EXTENSIONS = {
    "raw_result": {".csv", ".json", ".parquet", ".xlsx", ".mat", ".npz", ".npy"},
    "analysis_script": {".py", ".m", ".r", ".ipynb", ".jl"},
    "statistical_output": {".csv", ".json", ".xlsx", ".txt"},
    "table_or_figure": {".csv", ".xlsx", ".json", ".emf", ".pdf", ".png", ".svg", ".fig"},
    "model_checkpoint": {".pt", ".pth", ".ckpt", ".h5", ".keras", ".onnx"},
    "protocol": {".json", ".yaml", ".yml", ".toml", ".txt", ".md"},
}
# Claim-level raw evidence must expose searchable values and the experiment identifier.
# Binary archives can remain in artifacts.raw_results, but require a structured export
# before they can directly support a manuscript claim.
CLAIM_EVIDENCE_EXTENSIONS = {
    **ROLE_EXTENSIONS,
    "raw_result": {".csv", ".json", ".xlsx"},
    "statistical_output": {".csv", ".json", ".xlsx"},
}

METRIC_DIRECTION_WORDS = {
    "lower", "higher", "better", "worse", "smaller", "larger", "minimum", "maximum",
    "minimize", "maximize", "is", "than", "the", "and", "with", "without", "percent",
}


def normalized_identifier(value: Any) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", str(value).lower())).strip()


def primary_metric_identifier(value: Any) -> str:
    # The first clause is the declared metric; later clauses normally state direction.
    first_clause = re.split(r"[,;|]", str(value), maxsplit=1)[0]
    tokens = [token for token in re.findall(r"[a-z0-9]+", first_clause.lower()) if token not in METRIC_DIRECTION_WORDS]
    return " ".join(tokens)


def get(data: dict[str, Any], path: str, default: Any = None) -> Any:
    cur: Any = data
    for part in path.split("."):
        if not isinstance(cur, dict) or part not in cur:
            return default
        cur = cur[part]
    return cur


def populated(value: Any) -> bool:
    if value is None or value is False:
        return False
    if isinstance(value, str):
        stripped = value.strip().lower()
        return bool(stripped) and not stripped.startswith("replace with") and stripped not in {"path", "value", "identifier"}
    if isinstance(value, (list, dict)):
        return bool(value)
    return True


def missing_or_not_applicable(value: Any) -> bool:
    return not populated(value) or (isinstance(value, str) and value.strip().lower() == "not_applicable") or value is False


def descriptive_text(value: Any, minimum: int = 8) -> bool:
    if not isinstance(value, str) or len(value.strip()) < minimum:
        return False
    lowered = value.strip().lower()
    if any(token in lowered for token in ("not_applicable", "placeholder", "goes here", "todo", "tbd", "lorem ipsum", "generic", "something", "claimed to support")):
        return False
    words = re.findall(r"[a-zA-Z][a-zA-Z0-9_-]*", value)
    return len(words) >= 2 and len(set(word.lower() for word in words)) >= 2


def manuscript_locator(value: Any) -> bool:
    if isinstance(value, list):
        return bool(value) and all(isinstance(item, str) and len(item.strip()) >= 3 for item in value)
    return isinstance(value, str) and value.strip().lower().startswith("not_applicable:") and len(value.strip()) >= 20


def validate_evidence_entries(entries: Any, label: str, errors: list[str], expected_claim_id: str | None = None, expected_experiment_id: str | None = None) -> list[dict[str, Any]]:
    if not isinstance(entries, list) or not entries:
        errors.append(f"{label} must be a non-empty list of role/path objects.")
        return []
    valid: list[dict[str, Any]] = []
    for index, entry in enumerate(entries, 1):
        if not isinstance(entry, dict):
            errors.append(f"{label}[{index}] must be an object with role and path.")
            continue
        role, path = entry.get("role"), entry.get("path")
        if not isinstance(role, str) or role not in VALID_EVIDENCE_ROLES:
            errors.append(f"{label}[{index}].role must be one of {sorted(VALID_EVIDENCE_ROLES)}.")
            continue
        if not isinstance(path, str) or not path.strip():
            errors.append(f"{label}[{index}].path must be a non-empty path string.")
            continue
        extension = Path(path).suffix.lower()
        if extension not in CLAIM_EVIDENCE_EXTENSIONS[role]:
            errors.append(f"{label}[{index}] role {role} is incompatible with extension {extension or '<none>'}.")
            continue
        claim_id, experiment_id, purpose, markers = entry.get("claim_id"), entry.get("experiment_id"), entry.get("purpose"), entry.get("content_markers")
        if not isinstance(claim_id, str) or not claim_id.strip() or (expected_claim_id and claim_id != expected_claim_id):
            errors.append(f"{label}[{index}].claim_id must match {expected_claim_id or 'a declared claim'}.")
            continue
        if not isinstance(experiment_id, str) or experiment_id != expected_experiment_id:
            errors.append(f"{label}[{index}].experiment_id must match study.experiment_id.")
            continue
        if not isinstance(purpose, str) or len(purpose.strip()) < 15 or "placeholder" in purpose.lower() or "goes here" in purpose.lower():
            errors.append(f"{label}[{index}].purpose must specifically describe how the artifact supports the claim.")
            continue
        if not isinstance(markers, list) or not 1 <= len(markers) <= 6 or not all(isinstance(marker, str) and len(marker.strip()) >= 3 and not any(token in marker.lower() for token in ("generic", "placeholder", "something", "todo", "tbd")) for marker in markers):
            errors.append(f"{label}[{index}].content_markers must contain 1-6 specific identifiers expected in the artifact.")
            continue
        valid.append({"role": role, "path": path, "claim_id": claim_id, "experiment_id": experiment_id, "purpose": purpose, "content_markers": markers})
    return valid


def validate_manifest(data: dict[str, Any], paths_verified: bool = False) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    status = get(data, "study.status")
    if status not in VALID_STATUS:
        errors.append(f"study.status must be one of {sorted(VALID_STATUS)}.")
    experiment_id = get(data, "study.experiment_id")
    evidence_experiment_ids = {
        str(entry.get("experiment_id"))
        for claim in get(data, "study.claims", []) if isinstance(get(data, "study.claims", []), list) and isinstance(claim, dict)
        for entry in claim.get("evidence_artifacts", []) if isinstance(claim.get("evidence_artifacts"), list) and isinstance(entry, dict) and entry.get("experiment_id")
    }
    if experiment_id is None and status != "Submission candidate" and not paths_verified:
        experiment_id = next(iter(evidence_experiment_ids)) if len(evidence_experiment_ids) == 1 else "00000000-0000-0000-0000-000000000000"
    try:
        uuid.UUID(str(experiment_id))
    except (ValueError, TypeError, AttributeError):
        errors.append("study.experiment_id must be a valid UUID embedded in generated evidence artifacts.")

    required = [
        "study.title", "study.target_venue", "study.claims",
        "data.development_test_split", "data.sealed_test_split", "data.provenance_file",
        "methods.proposed", "methods.baselines",
        "statistics.primary_metrics", "statistics.evaluation_unit", "statistics.interval_method",
        "timing.hardware", "timing.software",
        "artifacts.raw_results", "artifacts.table_exports", "artifacts.figure_manifest",
    ]
    for path in required:
        if not populated(get(data, path)):
            errors.append(f"Missing or placeholder field: {path}")
    code_version = get(data, "artifacts.code_version", get(data, "artifacts.code_version_or_hash"))
    if not isinstance(code_version, str) or missing_or_not_applicable(code_version):
        errors.append("Missing or placeholder field: artifacts.code_version (legacy artifacts.code_version_or_hash is also accepted)")

    claims = get(data, "study.claims", [])
    if not isinstance(claims, list) or not claims:
        errors.append("study.claims must be a non-empty list.")
        claims = []
    else:
        for i, claim in enumerate(claims, 1):
            if not isinstance(claim, dict):
                errors.append(f"study.claims[{i}] must be an object.")
                continue
            for field in ("id", "statement", "claim_type", "primary_metric", "failure_condition", "test_contract", "evidence_artifacts"):
                if not populated(claim.get(field)):
                    errors.append(f"study.claims[{i}].{field} is missing or a placeholder.")
            if claim.get("claim_type") not in VALID_CLAIM_TYPES:
                errors.append(f"study.claims[{i}].claim_type must be one of {sorted(VALID_CLAIM_TYPES)}.")
            contract = claim.get("test_contract")
            if not isinstance(contract, dict):
                errors.append(f"study.claims[{i}].test_contract must be an object.")
            else:
                experiment_or_analysis_id = contract.get("experiment_or_analysis_id")
                if not isinstance(experiment_or_analysis_id, str) or len(experiment_or_analysis_id.strip()) < 3 or not populated(experiment_or_analysis_id):
                    errors.append(f"study.claims[{i}].test_contract.experiment_or_analysis_id requires a stable non-placeholder identifier.")
                for field in ("strongest_comparator", "evaluation_unit", "statistical_decision"):
                    if not descriptive_text(contract.get(field), minimum=5):
                        errors.append(f"study.claims[{i}].test_contract.{field} requires a specific descriptive value.")
                for field in ("figure_ids", "table_ids"):
                    if not manuscript_locator(contract.get(field)):
                        errors.append(
                            f"study.claims[{i}].test_contract.{field} must list manuscript IDs or use 'not_applicable:' with a reason."
                        )
                if isinstance(contract.get("figure_ids"), str) and isinstance(contract.get("table_ids"), str):
                    errors.append(f"study.claims[{i}].test_contract must map the claim to at least one Figure or Table.")
            validate_evidence_entries(claim.get("evidence_artifacts"), f"study.claims[{i}].evidence_artifacts", errors, expected_claim_id=str(claim.get("id", "")), expected_experiment_id=str(experiment_id))

    claim_types = {claim.get("claim_type") for claim in claims if isinstance(claim, dict)}
    for i, claim in enumerate(claims, 1):
        if not isinstance(claim, dict):
            continue
        entries = claim.get("evidence_artifacts", []) if isinstance(claim.get("evidence_artifacts"), list) else []
        roles = {entry.get("role") for entry in entries if isinstance(entry, dict)}
        claim_type = claim.get("claim_type")
        required_roles = {"raw_result"}
        if claim_type in {"performance", "model_comparison", "stability", "generalization", "ood_generalization", "composite_synergy"}:
            if not roles & {"analysis_script", "statistical_output"}:
                errors.append(f"study.claims[{i}] type {claim_type} requires an analysis_script or statistical_output artifact.")
        if claim_type == "predictive_uncertainty":
            required_roles |= {"analysis_script", "statistical_output"}
        if claim_type == "deployment":
            required_roles |= {"statistical_output"}
        for role in required_roles:
            if role not in roles:
                errors.append(f"study.claims[{i}] type {claim_type} requires a {role} artifact.")
        metric_id = primary_metric_identifier(claim.get("primary_metric", ""))
        marker_ids = {normalized_identifier(marker) for entry in entries if isinstance(entry, dict) for marker in entry.get("content_markers", []) if isinstance(marker, str)}
        if not metric_id or metric_id not in marker_ids:
            errors.append(f"study.claims[{i}] evidence content_markers must include the exact primary-metric identifier from its first clause.")
        if status == "Submission candidate" and claim_type in {
            "performance", "model_comparison", "stability", "generalization", "ood_generalization",
            "predictive_uncertainty", "composite_synergy",
        }:
            for role in ("analysis_script", "statistical_output"):
                if role not in roles:
                    errors.append(f"Submission claim study.claims[{i}] requires both reproducible analysis_script and structured statistical_output evidence; missing {role}.")

    baselines = get(data, "methods.baselines", [])
    if not isinstance(baselines, list) or not baselines:
        errors.append("methods.baselines must be a non-empty list.")
        baselines = []
    else:
        for i, baseline in enumerate(baselines, 1):
            if not isinstance(baseline, dict):
                errors.append(f"methods.baselines[{i}] must be an object.")
                continue
            for field in ("name", "role"):
                if not populated(baseline.get(field)):
                    errors.append(f"methods.baselines[{i}].{field} is missing.")
            for field in ("same_environment", "same_budget_and_constraints", "tuning_budget_disclosed"):
                if baseline.get(field) is not True:
                    warnings.append(f"Baseline {i} has not confirmed {field}=true.")

    task_structure = get(data, "data.task_structure")
    if task_structure not in VALID_TASK_STRUCTURES:
        errors.append(f"data.task_structure must be one of {sorted(VALID_TASK_STRUCTURES)}.")
    method_structure = get(data, "methods.method_structure")
    if method_structure not in VALID_METHOD_STRUCTURES:
        errors.append(f"methods.method_structure must be one of {sorted(VALID_METHOD_STRUCTURES)}.")

    if len(get(data, "statistics.primary_metrics", [])) > 2:
        warnings.append("More than two primary metrics are declared; confirm multiplicity and claim hierarchy.")
    if get(data, "statistics.independent_training_runs", 0) < 5:
        warnings.append("Fewer than five independent training runs are recorded; do not make a strong learning-stability claim without justification.")
    if "stability" in claim_types and get(data, "statistics.independent_training_runs", 0) < 5:
        errors.append("A stability claim requires at least five independent training runs or a separately justified non-learning stability design.")
    if "model_comparison" in claim_types:
        if get(data, "statistics.paired_design") is not True:
            errors.append("A model_comparison claim requires statistics.paired_design=true when methods share evaluation units.")
        if get(data, "statistics.effect_size_reported") is not True:
            errors.append("A model_comparison claim requires statistics.effect_size_reported=true.")

    if get(data, "execution.postprocessing_present") is True:
        for field in (
            "execution.postprocessed_metrics_recomputed",
            "execution.raw_and_postprocessed_feasibility_reported",
            "execution.fallback_defined",
            "execution.mission_level_fallback_rate_reported",
        ):
            if get(data, field) is not True:
                errors.append(f"Post-processing is present, but {field} is not true.")

    if get(data, "uncertainty.calibration_claimed") is True:
        for field in (
            "uncertainty.coverage_reported",
            "uncertainty.interval_width_reported",
            "uncertainty.nll_or_proper_score_reported",
            "uncertainty.subgroup_calibration_reported",
        ):
            if get(data, field) is not True:
                errors.append(f"Calibration is claimed, but {field} is not true.")

    if task_structure == "time_series":
        for path in ("data.temporal_protocol.preserves_time_order", "data.temporal_protocol.preprocessing_fitted_without_future_data"):
            if get(data, path) is not True:
                errors.append(f"Time-series protocol incomplete: {path} must be true.")
        for path in ("data.temporal_protocol.forecast_origin_and_horizon", "data.temporal_protocol.window_and_stride", "data.temporal_protocol.backtest_design"):
            if not descriptive_text(get(data, path), minimum=5):
                errors.append(f"Time-series protocol incomplete: {path} requires a descriptive value.")
        backtest = str(get(data, "data.temporal_protocol.backtest_design", "")).lower()
        if re.search(r"\b(?:random|shuffle|shuffled)\b", backtest):
            errors.append("Time-series backtest_design cannot use random or shuffled folds across chronology.")

    generalization_level = get(data, "data.generalization.claim_level", "not_applicable")
    if generalization_level not in VALID_GENERALIZATION_LEVELS:
        errors.append(f"data.generalization.claim_level must be one of {sorted(VALID_GENERALIZATION_LEVELS)}.")
    transfer_levels = {"leave_site_out", "cross_region", "ood_regime"}
    if "ood_generalization" in claim_types and generalization_level != "ood_regime":
        errors.append("An ood_generalization claim requires data.generalization.claim_level=ood_regime.")
    if "generalization" in claim_types and generalization_level not in transfer_levels:
        errors.append("A generalization claim requires leave_site_out, cross_region, or ood_regime evidence; use performance/stability for weaker scopes.")
    if generalization_level in transfer_levels:
        for path in (
            "data.generalization.held_out_factor",
            "data.generalization.shift_definition",
        ):
            if not populated(get(data, path)) or get(data, path) == "not_applicable":
                errors.append(f"Generalization claim {generalization_level} requires {path}.")
        if get(data, "data.generalization.development_exclusion_confirmed") is not True:
            errors.append(f"Generalization claim {generalization_level} requires development_exclusion_confirmed=true.")

    forecasting_required = task_structure == "time_series" and bool(claim_types & {"performance", "model_comparison", "stability", "generalization", "ood_generalization", "composite_synergy"})
    if forecasting_required and get(data, "forecasting_statistics.applicable") is not True:
        errors.append("Time-series comparison claims require forecasting_statistics.applicable=true.")
    if get(data, "forecasting_statistics.applicable") is True:
        if not descriptive_text(get(data, "forecasting_statistics.loss_differential_defined"), minimum=8):
            errors.append("Forecasting inference requires a descriptive loss_differential_defined value.")
        if get(data, "forecasting_statistics.serial_dependence_handled") is not True:
            errors.append("Forecasting inference requires serial_dependence_handled=true.")
        if not descriptive_text(get(data, "forecasting_statistics.horizon_adjustment_or_block_design"), minimum=8):
            errors.append("Forecasting inference requires a descriptive horizon_adjustment_or_block_design value.")
        if get(data, "forecasting_statistics.multiple_comparisons_handled") is not True:
            errors.append("Forecasting inference requires multiple_comparisons_handled=true.")

    if "predictive_uncertainty" in claim_types and get(data, "uncertainty.calibration_claimed") is not True:
        errors.append("A predictive_uncertainty claim requires uncertainty.calibration_claimed=true and calibration/sharpness evidence.")

    composite_required = "composite_synergy" in claim_types
    if composite_required and get(data, "composite_method.applicable") is not True:
        errors.append("A composite_synergy claim requires composite_method.applicable=true.")
    if method_structure in {"composite", "pipeline"} and not composite_required and get(data, "composite_method.applicable") is not True:
        warnings.append("Composite/pipeline method declared without a composite_synergy claim; component evidence is recommended but not a hard gate.")
    if get(data, "composite_method.applicable") is True:
        composite_claims = {str(claim.get("id")) for claim in claims if isinstance(claim, dict) and claim.get("claim_type") == "composite_synergy"}
        composite_claim_id = get(data, "composite_method.claim_id")
        if composite_claim_id is None and len(composite_claims) == 1:
            composite_claim_id = next(iter(composite_claims))
        if composite_claim_id not in composite_claims:
            errors.append("composite_method.claim_id must reference a declared composite_synergy claim.")
        factor_a_name = get(data, "composite_method.factor_a_name")
        factor_b_name = get(data, "composite_method.factor_b_name")
        if not descriptive_text(factor_a_name, minimum=5) or not descriptive_text(factor_b_name, minimum=5) or factor_a_name.strip().lower() == factor_b_name.strip().lower():
            errors.append("Composite-method evidence requires distinct descriptive factor_a_name and factor_b_name values.")
        factorial = get(data, "composite_method.factorial_variants", [])
        combinations: set[tuple[bool, bool]] = set()
        if not isinstance(factorial, list) or len(factorial) != 4:
            errors.append("Composite-method evidence requires exactly four factorial variant objects.")
        else:
            variant_names: list[str] = []
            for index, variant in enumerate(factorial, 1):
                if not isinstance(variant, dict) or not descriptive_text(variant.get("name"), minimum=5):
                    errors.append(f"composite_method.factorial_variants[{index}] requires a descriptive name.")
                    continue
                variant_names.append(variant["name"].strip().lower())
                lowered_name = variant["name"].strip().lower()
                if variant.get("factor_a") is True and str(factor_a_name).strip().lower() not in lowered_name:
                    errors.append(f"composite_method.factorial_variants[{index}] name must identify factor_a when factor_a=true.")
                if variant.get("factor_b") is True and str(factor_b_name).strip().lower() not in lowered_name:
                    errors.append(f"composite_method.factorial_variants[{index}] name must identify factor_b when factor_b=true.")
                if not isinstance(variant.get("factor_a"), bool) or not isinstance(variant.get("factor_b"), bool):
                    errors.append(f"composite_method.factorial_variants[{index}] requires boolean factor_a and factor_b.")
                    continue
                combinations.add((variant["factor_a"], variant["factor_b"]))
            if combinations != {(False, False), (True, False), (False, True), (True, True)}:
                errors.append("Composite factorial variants must cover all four factor_a/factor_b combinations exactly once.")
            if len(set(variant_names)) != len(variant_names):
                errors.append("Composite factorial variant names must be unique.")
        components = get(data, "composite_method.component_only_variants", [])
        if not isinstance(components, list) or len(components) < 2 or not all(descriptive_text(item, minimum=5) for item in components) or len({item.strip().lower() for item in components if isinstance(item, str)}) != len(components):
            errors.append("Composite-method evidence requires at least two named component-only variants.")
        if method_structure not in {"composite", "pipeline"}:
            errors.append("Composite-method evidence requires methods.method_structure to be composite or pipeline.")
        capacity = get(data, "composite_method.capacity_matched_baseline")
        if not isinstance(capacity, dict) or not descriptive_text(capacity.get("name"), minimum=5) or capacity.get("same_tuning_budget") is not True:
            errors.append("Composite-method capacity_matched_baseline requires a named baseline and same_tuning_budget=true.")
        elif not isinstance(capacity.get("parameter_difference_percent"), (int, float)) or isinstance(capacity.get("parameter_difference_percent"), bool):
            errors.append("Composite-method capacity_matched_baseline requires numeric parameter_difference_percent.")
        placebo = get(data, "composite_method.coupling_placebo_variant")
        if not isinstance(placebo, dict) or not descriptive_text(placebo.get("name"), minimum=5) or not descriptive_text(placebo.get("disrupted_path"), minimum=10) or placebo.get("capacity_preserved") is not True:
            errors.append("Composite-method coupling_placebo_variant requires a name, disrupted_path description, and capacity_preserved=true.")
        mechanism = get(data, "composite_method.coupling_mechanism")
        if not isinstance(mechanism, dict) or any(not descriptive_text(mechanism.get(field), minimum=8) for field in ("source_tensor", "operation", "destination", "claimed_new_capability")):
            errors.append("Composite-method coupling_mechanism requires source_tensor, operation, destination, and claimed_new_capability.")
        contrast = get(data, "composite_method.interaction_contrast")
        if not isinstance(contrast, dict) or any(not descriptive_text(contrast.get(field), minimum=5) for field in ("formula", "loss_direction", "evaluation_unit", "interval_method")):
            errors.append("Composite-method interaction_contrast requires formula, loss_direction, evaluation_unit, and interval_method.")
        signature = get(data, "composite_method.mechanism_signature")
        if not isinstance(signature, dict) or any(not descriptive_text(signature.get(field), minimum=5) for field in ("regime", "metric", "direction", "selection_source")):
            errors.append("Composite-method mechanism_signature requires regime, metric, direction, and selection_source.")
        tolerance = get(data, "composite_method.capacity_tolerance_percent")
        if not isinstance(tolerance, (int, float)) or isinstance(tolerance, bool) or not 0 <= tolerance <= 20:
            errors.append("Composite-method capacity_tolerance_percent must be numeric and between 0 and 20.")
        elif isinstance(capacity, dict) and isinstance(capacity.get("parameter_difference_percent"), (int, float)) and not isinstance(capacity.get("parameter_difference_percent"), bool) and abs(capacity["parameter_difference_percent"]) > tolerance:
            errors.append("Capacity-matched baseline parameter difference exceeds composite_method.capacity_tolerance_percent.")
        for field in ("interaction_interval_reported", "mechanism_signature_prespecified", "mechanism_signature_tested"):
            if get(data, f"composite_method.{field}") is not True:
                errors.append(f"Composite-method evidence requires composite_method.{field}=true.")
        artifacts = validate_evidence_entries(get(data, "composite_method.evidence_artifacts"), "composite_method.evidence_artifacts", errors, expected_experiment_id=str(experiment_id))
        for index, artifact in enumerate(artifacts, 1):
            if artifact["claim_id"] != composite_claim_id:
                errors.append(f"composite_method.evidence_artifacts[{index}].claim_id must match composite_method.claim_id.")
        roles = {entry["role"] for entry in artifacts}
        for required_role in ("raw_result", "statistical_output", "analysis_script"):
            if required_role not in roles:
                errors.append(f"Composite-method evidence requires a {required_role} artifact.")

    deployment_required = "deployment" in claim_types or get(data, "timing.deployment_claimed") is True
    if "deployment" in claim_types and get(data, "timing.deployment_claimed") is not True:
        errors.append("A deployment claim requires timing.deployment_claimed=true.")
    timing_fields = ("mean_reported", "p50_reported", "p95_reported", "end_to_end_cycle")
    if deployment_required:
        for field in timing_fields:
            if get(data, f"timing.{field}") is not True:
                errors.append(f"Deployment claim requires timing.{field}=true.")
    if deployment_required and get(data, "timing.end_to_end_cycle") is True:
        stages = get(data, "timing.pipeline_stages_included", [])
        online_stages = get(data, "methods.online_pipeline_stages", [])
        if not isinstance(online_stages, list) or len(online_stages) < 3 or len(set(online_stages)) != len(online_stages):
            errors.append("Deployment evidence requires at least three named methods.online_pipeline_stages.")
            online_stages = []
        if not isinstance(stages, list) or len(stages) < 3 or len(set(stages)) != len(stages):
            errors.append("End-to-end timing requires at least three timing.pipeline_stages_included entries.")
            stages = []
        missing_stages = sorted(set(online_stages) - set(stages))
        if missing_stages:
            errors.append(f"End-to-end timing omits required online pipeline stages: {missing_stages}.")
        categories = {
            "input": {"ingestion", "input", "data ingestion", "stream input"},
            "processing": {"preprocessing", "feature generation", "decomposition", "normalization"},
            "inference": {"inference", "model inference", "prediction"},
            "output": {"post-processing", "postprocessing", "output", "serialization"},
        }
        stage_tokens = {str(stage).strip().lower() for stage in stages}
        for category, aliases in categories.items():
            if not stage_tokens & aliases:
                errors.append(f"Deployment pipeline is missing the {category} stage category.")
        for field in ("p99_reported", "throughput_reported", "peak_memory_reported", "operational_deadline_compared"):
            if get(data, f"timing.{field}") is not True:
                errors.append(f"Deployment claim requires timing.{field}=true.")

    submission_checks = {
        "data.test_never_used_for_selection": True,
        "gates.no_placeholder_results": True,
        "gates.all_claims_traceable": True,
        "gates.negative_results_disclosed": True,
        "gates.all_figures_reproduce_saved_data": True,
        "gates.render_qa_passed": True,
    }
    if status == "Submission candidate":
        if not paths_verified:
            errors.append("Submission candidate requires physical artifact verification with --base-dir.")
        for path, expected in submission_checks.items():
            if get(data, path) is not expected:
                errors.append(f"Submission gate failed: {path} must be true.")
        for i, baseline in enumerate(baselines if isinstance(baselines, list) else [], 1):
            for field in ("same_environment", "same_budget_and_constraints", "tuning_budget_disclosed"):
                if baseline.get(field) is not True:
                    errors.append(f"Submission gate failed: methods.baselines[{i}].{field} must be true.")

    return {
        "status": status,
        "submission_candidate_ready": status == "Submission candidate" and not errors,
        "error_count": len(errors),
        "warning_count": len(warnings),
        "errors": errors,
        "warnings": warnings,
    }


def validate_paths(data: dict[str, Any], base_dir: Path, verify_digests: bool = False) -> list[str]:
    errors: list[str] = []
    candidates: list[tuple[str, Any, dict[str, Any] | None]] = [
        ("data.provenance_file", get(data, "data.provenance_file"), None),
        ("artifacts.figure_manifest", get(data, "artifacts.figure_manifest"), None),
    ]
    raw_results = get(data, "artifacts.raw_results", [])
    if not isinstance(raw_results, list) or not raw_results:
        errors.append("Path verification failed: artifacts.raw_results must be a non-empty list.")
        raw_results = []
    for path in raw_results:
        if not isinstance(path, str) or Path(path).suffix.lower() not in ROLE_EXTENSIONS["raw_result"]:
            errors.append(f"Path verification failed: artifacts.raw_results has an invalid path or extension: {path}")
        else:
            candidates.append(("artifacts.raw_results", path, None))
    table_exports = get(data, "artifacts.table_exports", [])
    if not isinstance(table_exports, list) or not table_exports:
        errors.append("Path verification failed: artifacts.table_exports must be a non-empty list.")
        table_exports = []
    for path in table_exports:
        if not isinstance(path, str) or Path(path).suffix.lower() not in {".csv", ".xlsx", ".json"}:
            errors.append(f"Path verification failed: artifacts.table_exports has an invalid path or extension: {path}")
        else:
            candidates.append(("artifacts.table_exports", path, None))
    # checkpoint_reference (or legacy checkpoint_hashes) is provenance metadata:
    # it may name a run/registry ID rather than a local file. Physical checkpoints
    # requiring path checks belong in evidence_artifacts with model_checkpoint role.
    for index, claim in enumerate(get(data, "study.claims", []) if isinstance(get(data, "study.claims", []), list) else [], 1):
        if isinstance(claim, dict):
            for entry in claim.get("evidence_artifacts", []) if isinstance(claim.get("evidence_artifacts"), list) else []:
                if isinstance(entry, dict):
                    candidates.append((f"study.claims[{index}].evidence_artifacts[{entry.get('role', 'unknown')}]", entry.get("path"), entry))
    for entry in get(data, "composite_method.evidence_artifacts", []) if isinstance(get(data, "composite_method.evidence_artifacts", []), list) else []:
        if isinstance(entry, dict):
            candidates.append((f"composite_method.evidence_artifacts[{entry.get('role', 'unknown')}]", entry.get("path"), entry))
    for label, value, evidence in candidates:
        if not isinstance(value, str) or not value.strip():
            errors.append(f"Path verification failed: {label} is not a path string.")
            continue
        candidate = Path(value)
        if not candidate.is_absolute():
            candidate = base_dir / candidate
        if not candidate.exists() or not candidate.is_file() or candidate.stat().st_size == 0:
            errors.append(f"Path verification failed: {label} is not a non-empty file: {candidate}")
            continue
        if evidence:
            if verify_digests and "sha256" in evidence:
                digest = evidence["sha256"]
                if not isinstance(digest, str) or not re.fullmatch(r"[0-9a-fA-F]{64}", digest):
                    errors.append(f"Path verification failed: {label}.sha256 must be a 64-character hexadecimal digest.")
                else:
                    # Explicit consumer: verify the caller's declared integrity contract.
                    hasher = hashlib.sha256()
                    with candidate.open("rb") as stream:
                        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                            hasher.update(chunk)
                    if hasher.hexdigest() != digest.lower():
                        errors.append(f"Path verification failed: {label} SHA-256 mismatch for {candidate}")
            role = evidence.get("role")
            suffix = candidate.suffix.lower()
            searchable = ""
            if suffix in {".csv", ".json", ".txt", ".py", ".m", ".r", ".jl", ".ipynb", ".yaml", ".yml", ".toml", ".md"}:
                searchable = candidate.read_text(encoding="utf-8", errors="replace").lower()
            elif suffix == ".xlsx" and zipfile.is_zipfile(candidate):
                with zipfile.ZipFile(candidate) as workbook:
                    searchable = " ".join(workbook.read(name).decode("utf-8", errors="replace") for name in workbook.namelist() if name.startswith("xl/") and name.endswith(".xml")).lower()
            markers = [str(marker).lower() for marker in evidence.get("content_markers", [])]
            markers.append(str(evidence.get("experiment_id", "")).lower())
            missing_markers = [marker for marker in markers if marker not in searchable] if searchable else []
            if markers and searchable and missing_markers:
                errors.append(f"Evidence-content verification failed: {label} is missing declared markers {missing_markers}: {candidate}")
                continue
            if role == "raw_result" and suffix == ".csv":
                with candidate.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
                    rows = list(csv.reader(stream))
                numeric_cells = []
                for row in rows[1:]:
                    for cell in row:
                        try:
                            parsed = float(cell)
                            if math.isfinite(parsed):
                                numeric_cells.append(parsed)
                        except (TypeError, ValueError):
                            pass
                if len(rows) < 2 or len(rows[0]) < 1 or len(numeric_cells) < 2:
                    errors.append(f"Evidence-role verification failed: raw_result CSV requires at least two numeric observations: {candidate}")
                header_id = normalized_identifier(" ".join(rows[0])) if rows else ""
                declared_markers = [normalized_identifier(marker) for marker in evidence.get("content_markers", [])]
                if declared_markers and not any(marker and marker in header_id for marker in declared_markers):
                    errors.append(f"Evidence-role verification failed: raw_result CSV headers omit every declared content marker: {candidate}")
            elif role == "raw_result" and suffix == ".json":
                try:
                    payload = json.loads(candidate.read_text(encoding="utf-8"))
                except json.JSONDecodeError:
                    errors.append(f"Evidence-role verification failed: invalid raw_result JSON: {candidate}")
                else:
                    def numeric_values(value: Any) -> list[float]:
                        if isinstance(value, bool):
                            return []
                        if isinstance(value, (int, float)):
                            parsed = float(value)
                            return [parsed] if math.isfinite(parsed) else []
                        if isinstance(value, list):
                            return [number for item in value for number in numeric_values(item)]
                        if isinstance(value, dict):
                            return [number for item in value.values() for number in numeric_values(item)]
                        return []
                    numbers = numeric_values(payload)
                    if len(numbers) < 2:
                        errors.append(f"Evidence-role verification failed: raw_result JSON requires at least two numeric observations: {candidate}")
            elif role == "raw_result" and suffix == ".npy":
                try:
                    import numpy as np
                    array = np.load(candidate, allow_pickle=False)
                    if array.size < 2 or not np.issubdtype(array.dtype, np.number):
                        raise ValueError("array is non-numeric or has fewer than two observations")
                except Exception as exc:
                    errors.append(f"Evidence-role verification failed: invalid numeric NPY data ({exc}): {candidate}")
            elif role == "raw_result" and suffix == ".npz":
                try:
                    import numpy as np
                    with np.load(candidate, allow_pickle=False) as archive:
                        arrays = [archive[name] for name in archive.files]
                    if not arrays or sum(array.size for array in arrays if np.issubdtype(array.dtype, np.number)) < 2:
                        raise ValueError("archive lacks numeric observations")
                except Exception as exc:
                    errors.append(f"Evidence-role verification failed: invalid numeric NPZ data ({exc}): {candidate}")
            elif role == "raw_result" and suffix == ".xlsx":
                try:
                    from openpyxl import load_workbook
                    workbook = load_workbook(candidate, read_only=True, data_only=True)
                    values = [cell for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True) for cell in row]
                    workbook.close()
                    numbers = [value for value in values if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))]
                    if len(numbers) < 2:
                        raise ValueError("workbook lacks numeric observations")
                except Exception as exc:
                    errors.append(f"Evidence-role verification failed: invalid numeric XLSX data ({exc}): {candidate}")
            elif role == "raw_result" and suffix == ".parquet":
                payload = candidate.read_bytes()
                if len(payload) < 8 or not (payload.startswith(b"PAR1") and payload.endswith(b"PAR1")):
                    errors.append(f"Evidence-role verification failed: invalid Parquet signature: {candidate}")
            elif role == "raw_result" and suffix == ".mat":
                payload = candidate.read_bytes()[:128]
                if not (payload.startswith(b"MATLAB") or payload.startswith(b"\x89HDF\r\n\x1a\n")):
                    errors.append(f"Evidence-role verification failed: invalid MAT signature: {candidate}")
            elif role == "analysis_script":
                text = candidate.read_text(encoding="utf-8", errors="replace")
                if len(text.splitlines()) < 3 or not re.search(r"\b(?:def|function|read_csv|load|bootstrap|confidence|interval|mean|metric)\b", text, re.IGNORECASE):
                    errors.append(f"Evidence-role verification failed: analysis_script lacks executable analysis structure: {candidate}")
            elif role == "statistical_output" and suffix == ".json":
                try:
                    payload = json.loads(candidate.read_text(encoding="utf-8"))
                except json.JSONDecodeError:
                    errors.append(f"Evidence-role verification failed: invalid statistical_output JSON: {candidate}")
                else:
                    if not isinstance(payload, dict):
                        errors.append(f"Evidence-role verification failed: statistical_output JSON must be an object: {candidate}")
                        payload = {}
                    numeric_scalars = {str(key).lower(): value for key, value in payload.items() if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))}
                    intervals = [value for key, value in payload.items() if str(key).lower() in {"ci", "confidence_interval", "interval"} and isinstance(value, list) and len(value) == 2 and all(isinstance(item, (int, float)) and not isinstance(item, bool) and math.isfinite(float(item)) for item in value)]
                    has_estimate = bool(set(numeric_scalars) & {"estimate", "effect", "p_value", "lower", "upper"})
                    if not has_estimate or not (intervals or {"lower", "upper"} <= set(numeric_scalars) or "p_value" in numeric_scalars):
                        errors.append(f"Evidence-role verification failed: statistical_output JSON lacks estimate/interval fields: {candidate}")
            elif role == "statistical_output" and suffix == ".csv":
                with candidate.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
                    rows = list(csv.reader(stream))
                headers = {cell.strip().lower() for cell in rows[0]} if rows else set()
                recognized = {"estimate", "effect", "ci", "confidence_interval", "p_value", "lower", "upper"}
                recognized_indices = [index for index, cell in enumerate(rows[0] if rows else []) if cell.strip().lower() in recognized]
                numeric_estimates = 0
                for row in rows[1:]:
                    for index in recognized_indices:
                        if index < len(row):
                            try:
                                parsed = float(row[index])
                                if math.isfinite(parsed):
                                    numeric_estimates += 1
                            except (TypeError, ValueError):
                                pass
                if len(rows) < 2 or not headers & recognized or numeric_estimates < 1:
                    errors.append(f"Evidence-role verification failed: statistical_output CSV lacks estimate/interval columns and data: {candidate}")
            elif role == "statistical_output" and suffix == ".txt":
                text = candidate.read_text(encoding="utf-8", errors="replace")
                if not re.search(r"\b(?:estimate|effect|confidence interval|95% ci|p[-_ ]?value|lower|upper)\b", text, re.IGNORECASE) or not re.search(r"[-+]?\d+(?:\.\d+)?", text):
                    errors.append(f"Evidence-role verification failed: statistical_output TXT lacks labeled numerical estimates: {candidate}")
            elif role == "statistical_output" and suffix == ".xlsx":
                try:
                    from openpyxl import load_workbook
                    workbook = load_workbook(candidate, read_only=True, data_only=True)
                    rows = [row for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)]
                    workbook.close()
                    headers = {str(cell).strip().lower() for row in rows[:3] for cell in row if cell is not None}
                    numbers = [cell for row in rows[1:] for cell in row if isinstance(cell, (int, float)) and not isinstance(cell, bool) and math.isfinite(float(cell))]
                    if not numbers or not headers & {"estimate", "effect", "ci", "confidence_interval", "p_value", "lower", "upper"}:
                        raise ValueError("workbook lacks labeled numerical estimate/interval output")
                except Exception as exc:
                    errors.append(f"Evidence-role verification failed: invalid statistical_output XLSX ({exc}): {candidate}")
    return errors


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate a claim-to-evidence experiment manifest.")
    parser.add_argument("manifest")
    parser.add_argument("--json-out", default="")
    parser.add_argument("--base-dir", default="", help="Project root used to verify provenance, result, table, figure-manifest, and checkpoint files.")
    parser.add_argument("--verify-digests", action="store_true", help="Verify only declared legacy sha256 values; requires --base-dir. Does not create digests or sidecars.")
    args = parser.parse_args()
    if args.verify_digests and not args.base_dir:
        parser.error("--verify-digests requires --base-dir")
    path = Path(args.manifest).resolve()
    data = json.loads(path.read_text(encoding="utf-8"))
    paths_verified = bool(args.base_dir)
    result = {"manifest": str(path), **validate_manifest(data, paths_verified=paths_verified)}
    if paths_verified:
        path_errors = validate_paths(data, Path(args.base_dir).resolve(), verify_digests=args.verify_digests)
        result["errors"].extend(path_errors)
        result["error_count"] = len(result["errors"])
        result["submission_candidate_ready"] = result["status"] == "Submission candidate" and not result["errors"]
    payload = json.dumps(result, indent=2, ensure_ascii=False)
    sys.stdout.buffer.write((payload + "\n").encode("utf-8", errors="replace"))
    if args.json_out:
        Path(args.json_out).write_text(payload, encoding="utf-8")
    raise SystemExit(1 if result["error_count"] else 0)


if __name__ == "__main__":
    main()
