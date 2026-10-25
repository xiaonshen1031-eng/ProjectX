#!/usr/bin/env python3
"""Project delivery audit with shared final prose and Project evidence gates."""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

from docx import Document

from mathtype_authenticity import audit_mathtype
from audit_docx_submission import audit as audit_docx
from check_writing_quality import inspect_writing
from manuscript_prose_rules import read_reviews
from project_plot_policy import audit as audit_plot_policy


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def visual_review_matches_pdf(review: dict, pdf_path: Path) -> bool:
    declared = str(review.get("pdf_sha256", "")).lower()
    return bool(declared) and declared == file_hash(pdf_path).lower()


def audit(args) -> dict:
    manifest = load(args.manifest)
    errors, warnings = [], []
    if manifest.get("document_type") != "project" or manifest.get("manuscript_mode") != "project_word":
        errors.append("Workspace is not a project_word route")
    if manifest.get("language") not in {"chinese", "english"}:
        errors.append("Project language is missing")
    for path in args.data_audit:
        record = load(path)
        if record.get("status") != "pass": errors.append(f"Data audit failed: {path}")
    result, validation, figures = load(args.result), load(args.validation), load(args.figures)
    if result.get("status") != "optimal": errors.append("Declared solver result is not optimal")
    if validation.get("status") != "pass": errors.append("Independent validation did not pass")
    plot_policy = audit_plot_policy(figures, args.figures.parent)
    errors.extend(plot_policy["errors"])
    roles = set()
    for figure in figures.get("artifacts", []):
        path = Path(figure.get("path", ""))
        if not path.is_file() or path.stat().st_size == 0: errors.append(f"Missing figure: {path}")
        role = figure.get("role", "")
        if not role: errors.append(f"Figure lacks a technical role: {figure.get('id')}")
        if role in roles: warnings.append(f"Repeated figure role requires contextual review: {role}")
        roles.add(role)
        typography = figure.get("final_size_typography", {})
        if not (8 <= typography.get("tick_pt", 0) <= 10 and 8 <= typography.get("legend_pt", 0) <= 10 and 9 <= typography.get("axis_label_pt", 0) <= 11):
            errors.append(f"Figure final-size typography outside Project default: {figure.get('id')}")
    writing = None
    if not args.docx.is_file(): errors.append("Final DOCX is missing")
    else:
        doc = Document(args.docx)
        if not doc.paragraphs: errors.append("Final DOCX contains no visible paragraphs")
        text = "\n".join(p.text for p in doc.paragraphs)
        if "[按当前项目材料填写]" in text: errors.append("Template placeholder remains in final DOCX")
        if any(token in text for token in (r"\[", r"\]")): errors.append("ToggleTeX source remains in final DOCX")
        try:
            writing = inspect_writing(
                args.docx, manuscript_mode="project",
                language=manifest.get("language") or "auto",
                review_decisions=read_reviews(getattr(args, "review_decisions", None)),
            )
            if (writing["blocking_warning_count"] or writing["unresolved_editorial_count"]
                    or writing["editorial_review_errors"]):
                errors.append("Shared final language gate did not pass for the current Project DOCX")
        except (OSError, ValueError, TypeError) as exc:
            errors.append(f"Cannot complete current Project language review: {exc}")
    equation = None
    office_math = None
    backend = manifest.get("equation_backend", "mathtype_native")
    if backend == "mathtype_native":
        if not args.mathtype_evidence: errors.append("MathType authenticity evidence is required")
        elif args.docx.is_file():
            equation = audit_mathtype(args.docx, args.mathtype_evidence)
            if equation.get("status") != "pass": errors.append("Native MathType authenticity gate did not pass")
    elif backend == "office_math" and args.docx.is_file():
        office_math = audit_docx(args.docx, False, [], equation_backend="office_math", equations_required=True)
        if office_math.get("office_math_object_count", 0) == 0:
            errors.append("Office Math backend contains no editable OMML objects")
        if office_math.get("equation_backend_detected") != "office_math":
            errors.append("Office Math backend contains a missing or mixed equation representation")
        if office_math.get("toggle_tex_formula_count", 0):
            errors.append("ToggleTeX source remains in the Office Math delivery")
    else:
        errors.append(f"Unsupported Project equation backend: {backend}")
    if not args.pdf.is_file() or args.pdf.stat().st_size == 0:
        errors.append("Final PDF is missing")
        page_count = 0
    else:
        try:
            from pypdf import PdfReader
            page_count = len(PdfReader(str(args.pdf)).pages)
        except Exception as exc:
            errors.append(f"Cannot inspect PDF pages: {exc}"); page_count = 0
    review = load(args.visual_review)
    if args.pdf.is_file() and not visual_review_matches_pdf(review, args.pdf):
        errors.append("Visual review is not bound to the current PDF SHA-256")
    reviewed = sorted(set(int(v) for v in review.get("reviewed_pages", [])))
    if review.get("reviewer") in {None, "", "generator"}:
        errors.append("Visual review must identify the inspecting agent or person")
    if reviewed != list(range(1, page_count + 1)):
        errors.append("Visual review does not cover every PDF page")
    screenshots = [Path(v) for v in review.get("page_images", [])]
    if len(screenshots) != page_count or any(not p.is_file() for p in screenshots):
        errors.append("Visual review page images are incomplete")
    if review.get("status") != "pass": errors.append("Visual page review did not pass")
    return {
        "schema": "senior_projectx.project.delivery_audit.v1", "status": "pass" if not errors else "fail",
        "document_type": "project", "page_count": page_count, "errors": errors, "warnings": warnings,
        "mathtype_authenticity": equation,
        "office_math_audit": office_math,
        "writing_audit": writing,
        "scope": "Shared final prose requirements plus Project task/data/result/figure/formula/document delivery; no conference anonymity, journal highlights, or reviewer-score requirements.",
    }


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--manifest", required=True, type=Path); p.add_argument("--data-audit", action="append", required=True, type=Path)
    p.add_argument("--result", required=True, type=Path); p.add_argument("--validation", required=True, type=Path)
    p.add_argument("--figures", required=True, type=Path); p.add_argument("--docx", required=True, type=Path)
    p.add_argument("--pdf", required=True, type=Path); p.add_argument("--visual-review", required=True, type=Path)
    p.add_argument("--mathtype-evidence", type=Path); p.add_argument("--json-out", type=Path)
    p.add_argument("--review-decisions", type=Path, help="Current-text-bound scientific-content decisions for the shared final language gate; the current DOCX is always screened.")
    args = p.parse_args(argv); result = audit(args); payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out: args.json_out.write_text(payload, encoding="utf-8")
    sys.stdout.buffer.write((payload + "\n").encode("utf-8")); return 0 if result["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
